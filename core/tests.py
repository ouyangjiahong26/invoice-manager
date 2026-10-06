import base64
import json
import re
import shutil
import tempfile
import zipfile
from decimal import Decimal
from io import BytesIO, StringIO
from email.header import decode_header
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from django.contrib import admin
from django.contrib.auth.models import User
from django.contrib.messages.storage.cookie import CookieStorage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings

from openpyxl import load_workbook

import pymupdf

from . import pairing, staging, stamping, vision
from .audit import ActorMiddleware
from .models import Attachment, AuditLog, Batch, Category, Item
from .validation import check_item

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

TEMP_MEDIA_ROOT = tempfile.mkdtemp()

ALIPAY_ORDER_NO = "20" + "1" * 26  # 28 位，20 开头


def _png_upload(name):
    return SimpleUploadedFile(name, PNG_1X1, content_type="image/png")


def _content_disposition(response):
    """Django 把非 ASCII 响应头按 RFC 2047 编码，测试断言前先解码。"""
    return "".join(
        part.decode(charset or "utf-8")
        for part, charset in decode_header(response["Content-Disposition"])
    )


def _item_data(**overrides):
    data = {
        "title": "硬盘",
        "category": None,  # 由 setUp 填充
        "actual_amount": "100.00",
        "invoice_amount": "100.00",
        "invoices": [_png_upload("invoice.png")],
        "payments": [_png_upload("payment.png")],
    }
    data.update(overrides)
    return data


class SubmissionTestCase(TestCase):
    """条目提交、附件、校验、导出的公共基建。"""

    def setUp(self):
        # 清空媒体目录：文件不随 TestCase 事务回滚，跨用例残留会让重传同名
        # 文件获得随机后缀，破坏快照的确定性断言。
        shutil.rmtree(TEMP_MEDIA_ROOT, ignore_errors=True)
        self.addCleanup(shutil.rmtree, TEMP_MEDIA_ROOT, ignore_errors=True)
        self.category, _ = Category.objects.get_or_create(
            name="信息服务费", defaults={"order": 0}
        )
        self.batch = Batch.objects.create(name="2025 报销批次")
        self.student = User.objects.create_user(
            "zhangsan22", password="pw", first_name="张三"
        )
        self.student_b = User.objects.create_user(
            "lisi23", password="pw", first_name="李四"
        )
        self.staff = User.objects.create_user(
            "admin22", password="pw", first_name="管理员", is_staff=True
        )
        self.factory = RequestFactory()

    def _create_item(self, user, **overrides):
        self.client.force_login(user)
        overrides.setdefault("category", self.category.pk)
        data = _item_data(**overrides)
        if user.is_staff:
            data["owner"] = user.pk   # staff 表单含付款人字段，需显式提供
            data["status"] = Item.STATUS_PENDING  # staff 表单含状态字段，需显式提供
        response = self.client.post("/items/new/", data)
        self.assertRedirects(response, "/")
        return Item.objects.latest("id")

    def _stage(self, uploads):
        """走 stage 端点建会话，返回 (session_id, [file_id])。"""
        self.client.force_login(self.student)
        response = self.client.post("/items/batch/stage/", {"files": uploads})
        data = response.json()
        return data["session_id"], [f["id"] for f in data["files"]]



@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class AttachmentSubmissionTests(SubmissionTestCase):
    def test_create_multi_attachment_records_amounts(self):
        response = self.client.force_login(self.student)
        data = _item_data(
            category=self.category.pk,
            actual_amount="148.40",
            invoice_amount="148.40",
            invoices=[_png_upload("a.pdf"), _png_upload("b.pdf")],
            payments=[_png_upload("p.pdf")],
        )
        data["new_invoice_0_amount"] = "67.70"
        data["new_invoice_1_amount"] = "80.70"
        data["new_payment_0_amount"] = "148.40"
        data["new_payment_0_order_no"] = ALIPAY_ORDER_NO
        data["new_invoice_0_ocr"] = '{"buyer_name": "清华大学"}'
        response = self.client.post("/items/new/", data)
        self.assertRedirects(response, "/")

        item = Item.objects.latest("id")
        self.assertEqual(item.attachments.count(), 3)
        self.assertEqual(item.attachments.filter(kind=Attachment.KIND_INVOICE).count(), 2)
        first = item.attachments.filter(kind=Attachment.KIND_INVOICE).order_by("id").first()
        self.assertEqual(str(first.amount), "67.70")
        self.assertEqual(first.ocr_data, {"buyer_name": "清华大学"})
        self.assertEqual(item.attachments.get(kind=Attachment.KIND_PAYMENT).order_no, ALIPAY_ORDER_NO)

    def test_create_support_upload(self):
        """侧边栏"证明材料"上传块：supports 文件 + new_support 前缀字段。"""
        self.client.force_login(self.student)
        data = _item_data(
            category=self.category.pk,
            actual_amount="656.13",
            invoice_amount="656.13",
            invoices=[_png_upload("inv.pdf")],
            payments=[_png_upload("pay.pdf")],
        )
        data["supports"] = [_png_upload("flow.pdf")]
        data["new_invoice_0_amount"] = "656.13"
        data["new_payment_0_amount"] = "1096.50"
        data["new_payment_0_order_no"] = ALIPAY_ORDER_NO
        data["new_support_0_amount"] = "1096.50"
        response = self.client.post("/items/new/", data)
        self.assertRedirects(response, "/")
        item = Item.objects.latest("id")
        support = item.attachments.get(kind=Attachment.KIND_SUPPORT)
        self.assertEqual(str(support.amount), "1096.50")

    def test_create_warns_on_amount_mismatch(self):
        self.client.force_login(self.student)
        data = _item_data(
            category=self.category.pk,
            actual_amount="100.00",
            invoices=[_png_upload("a.png")],
            payments=[_png_upload("p.png")],
        )
        data["new_payment_0_amount"] = "90.00"
        data["new_payment_0_order_no"] = ALIPAY_ORDER_NO
        response = self.client.post("/items/new/", data, follow=True)
        self.assertRedirects(response, "/")
        self.assertContains(response, "支付与退款附件合计与填写的实付款不一致，请核对")

    def test_create_matching_amounts_has_no_amount_warning(self):
        self.client.force_login(self.student)
        data = _item_data(
            category=self.category.pk,
            actual_amount="179.00",
            invoice_amount="199.00",
            invoices=[_png_upload("a.png")],
            payments=[_png_upload("p.png")],
            refunds=[_png_upload("r.png")],
        )
        data["new_invoice_0_amount"] = "199.00"
        data["new_payment_0_amount"] = "199.00"
        data["new_payment_0_order_no"] = ALIPAY_ORDER_NO
        data["new_refund_0_amount"] = "20.00"
        response = self.client.post("/items/new/", data, follow=True)
        self.assertRedirects(response, "/")
        self.assertNotContains(response, "支付与退款附件合计与填写的实付款不一致")
        self.assertNotContains(response, "发票附件金额合计与填写的发票金额不一致")

    def test_non_dict_ocr_is_dropped(self):
        self.client.force_login(self.student)
        data = _item_data(category=self.category.pk, payments=[])
        data["new_invoice_0_amount"] = "100.00"
        data["new_invoice_0_ocr"] = "5"
        response = self.client.post("/items/new/", data, follow=True)
        self.assertRedirects(response, "/")
        invoice = Item.objects.latest("id").attachments.get(kind=Attachment.KIND_INVOICE)
        self.assertIsNone(invoice.ocr_data)

    def test_rejects_disallowed_extension(self):
        self.client.force_login(self.student)
        data = _item_data(
            category=self.category.pk,
            invoices=[SimpleUploadedFile("evil.exe", b"x")],
            payments=[],
        )
        response = self.client.post("/items/new/", data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "仅支持 jpg / jpeg / png / pdf 文件")
        self.assertFalse(Item.objects.exists())

    def test_attachment_delete_by_other_user_forbidden(self):
        item = self._create_item(self.student)
        attachment = item.attachments.first()
        self.client.force_login(self.student_b)
        response = self.client.post(f"/items/attachments/{attachment.pk}/delete/")
        self.assertEqual(response.status_code, 403)
        self.assertTrue(Attachment.objects.filter(pk=attachment.pk).exists())

    def test_attachment_update_edits_amount(self):
        item = self._create_item(self.student)
        attachment = item.attachments.get(kind=Attachment.KIND_PAYMENT)
        self.client.force_login(self.student)
        response = self.client.post(
            f"/items/attachments/{attachment.pk}/edit/",
            {"kind": Attachment.KIND_PAYMENT, "amount": "55.50",
             "order_no": ALIPAY_ORDER_NO, "merchant_no": "", "invoice_no": ""},
        )
        self.assertRedirects(response, "/")
        attachment.refresh_from_db()
        self.assertEqual(str(attachment.amount), "55.50")

    def test_attachment_delete_by_owner_removes_record(self):
        item = self._create_item(self.student)
        attachment = item.attachments.get(kind=Attachment.KIND_INVOICE)
        self.client.force_login(self.student)
        response = self.client.post(f"/items/attachments/{attachment.pk}/delete/")
        self.assertRedirects(response, "/")
        self.assertFalse(Attachment.objects.filter(pk=attachment.pk).exists())
        self.assertTrue(item.attachments.filter(kind=Attachment.KIND_PAYMENT).exists())

    def test_file_serve_requires_login(self):
        item = self._create_item(self.student)
        attachment = item.attachments.first()
        self.client.logout()
        response = self.client.get(f"/items/attachments/{attachment.pk}/file/")
        self.assertEqual(response.status_code, 302)

    def test_prefill_view_returns_vision_payload(self):
        self.client.force_login(self.student)
        payload = {"amount": 19.9, "platform": "alipay", "order_no": ALIPAY_ORDER_NO,
                   "merchant_no": "", "_source": "image"}
        with mock.patch("core.views.vision_prefill", return_value=payload) as vision:
            response = self.client.post(
                "/items/prefill/", {"file": _png_upload("p.png"), "kind": "payment"}
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), payload)
        self.assertEqual(vision.call_args[0][2], "payment")

    def test_prefill_view_rejects_bad_kind_and_oversize(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/prefill/", {"file": _png_upload("p.png"), "kind": "unknown"}
        )
        self.assertEqual(response.status_code, 400)
        big = SimpleUploadedFile("big.png", b"0" * (10 * 1024 * 1024 + 1),
                                content_type="image/png")
        response = self.client.post("/items/prefill/", {"file": big, "kind": "invoice"})
        self.assertEqual(response.status_code, 400)

    def test_export_excel_header_and_hyperlinks(self):
        item = self._create_item(self.student)
        item.status = Item.STATUS_APPROVED
        item.save()
        self.client.force_login(self.staff)
        response = self.client.get("/export/excel/", {"batch": self.batch.pk})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["序号", "明细", "类别", "实付款", "付款人", "支付记录", "票据"],
        )
        self.assertIn(f"/items/attachments/{item.first_payment.pk}/file/",
                      sheet.cell(row=2, column=6).hyperlink.target)

    def test_export_zip_names_include_kind_labels(self):
        item = self._create_item(self.student)
        item.status = Item.STATUS_APPROVED
        item.save()
        self.client.force_login(self.staff)
        response = self.client.get("/export/zip/", {"batch": self.batch.pk})
        names = zipfile.ZipFile(BytesIO(response.content)).namelist()
        self.assertEqual(len(names), 2)
        self.assertTrue(any("_发票_1." in name for name in names))
        self.assertTrue(any("_支付记录_1." in name for name in names))


class OrderLineTests(SimpleTestCase):
    """发票首页单号行前缀称谓：platform 优先、单号前缀推断、兜底。空单号跳过。"""

    @staticmethod
    def _payment(order_no, ocr_data=None):
        return SimpleNamespace(order_no=order_no, ocr_data=ocr_data)

    def test_label_prefers_platform_then_prefix_then_fallback(self):
        wechat = "4" + "1" * 27
        self.assertEqual(
            stamping.order_lines([self._payment(ALIPAY_ORDER_NO, {"platform": "wechat"})]),
            [f"交易单号: {ALIPAY_ORDER_NO}"],  # platform 优先于单号前缀
        )
        self.assertEqual(stamping.order_lines([self._payment(ALIPAY_ORDER_NO)]),
                         [f"订单号: {ALIPAY_ORDER_NO}"])
        self.assertEqual(stamping.order_lines([self._payment(wechat)]),
                         [f"交易单号: {wechat}"])
        self.assertEqual(
            stamping.order_lines([self._payment("X123", {"platform": "other"})]),
            ["平台单号: X123"],
        )

    def test_label_jd_order_no(self):
        """京东订单编号判不出平台（other）时按前缀称总订单编号。"""
        jd = "34" + "1" * 14
        self.assertEqual(
            stamping.order_lines([self._payment(jd, {"platform": "other"})]),
            [f"总订单编号: {jd}"],
        )

    def test_blank_order_no_skipped(self):
        self.assertEqual(
            stamping.order_lines([self._payment(""), self._payment(None)]), []
        )


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ExportStampingTests(SubmissionTestCase):
    """导出 zip 票据标注：全部附件每页盖序号、第一张发票附件首页写单号、图片转 PDF。"""

    @staticmethod
    def _pdf_bytes(pages=1):
        doc = pymupdf.open()
        for _ in range(pages):
            doc.new_page()
        return doc.tobytes()

    @staticmethod
    def _png_bytes(width=200, height=200):
        """生成空白 PNG。1x1 图转 PDF 页高不足 1 pt，标注文字会落在页外。"""
        return pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height)).tobytes("png")

    def _item(self, **kwargs):
        data = {"owner": self.student, "title": "书", "category": self.category,
                "batch": self.batch, "actual_amount": "100.00",
                "invoice_amount": "100.00", "status": Item.STATUS_APPROVED}
        data.update(kwargs)
        return Item.objects.create(**data)

    def _attach(self, item, kind, name, data, **kwargs):
        return Attachment.objects.create(
            item=item, kind=kind, file=SimpleUploadedFile(name, data), **kwargs)

    def _export(self):
        self.client.force_login(self.staff)
        response = self.client.get("/export/zip/", {"batch": self.batch.pk})
        return zipfile.ZipFile(BytesIO(response.content))

    @staticmethod
    def _red_spans(data):
        """把 PDF 字节解析成 [(页号, 红色文本)]，只取标注色文字，不依赖 span 顺序。"""
        doc = pymupdf.open(stream=data, filetype="pdf")
        return [
            (pno, span["text"])
            for pno, page in enumerate(doc)
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
            if span["color"] == 0xFF0000
        ]

    def test_stamps_every_page_and_order_lines_on_first_invoice(self):
        item = self._item()
        self._attach(item, Attachment.KIND_INVOICE, "inv.pdf", self._pdf_bytes(2))
        self._attach(item, Attachment.KIND_INVOICE, "support.pdf", self._pdf_bytes(1))
        self._attach(item, Attachment.KIND_PAYMENT, "pay.png", self._png_bytes(),
                     order_no=ALIPAY_ORDER_NO, ocr_data={"platform": "alipay"})
        self._attach(item, Attachment.KIND_REFUND, "ref.pdf", self._pdf_bytes(1))
        zf = self._export()
        names = zf.namelist()
        self.assertEqual(len(names), 4)
        invoice, support = (n for n in sorted(names) if "_发票_" in n)
        payment = next(n for n in names if "_支付记录_" in n)
        refund = next(n for n in names if "_退款记录_" in n)
        self.assertTrue(payment.endswith(".pdf"))  # 图片附件导出为 PDF
        self.assertEqual(
            sorted(self._red_spans(zf.read(invoice))),
            sorted([(0, "01"), (0, f"订单号: {ALIPAY_ORDER_NO}"), (1, "01")]),
        )
        self.assertEqual(self._red_spans(zf.read(support)), [(0, "01")])  # 佐证只盖序号
        self.assertEqual(self._red_spans(zf.read(payment)), [(0, "01")])
        self.assertEqual(self._red_spans(zf.read(refund)), [(0, "01")])

    def test_multiple_payments_all_written(self):
        item = self._item()
        self._attach(item, Attachment.KIND_INVOICE, "inv.pdf", self._pdf_bytes())
        wechat = "4" + "1" * 27
        self._attach(item, Attachment.KIND_PAYMENT, "p1.png", self._png_bytes(),
                     order_no=ALIPAY_ORDER_NO)
        self._attach(item, Attachment.KIND_PAYMENT, "p2.png", self._png_bytes(),
                     order_no=wechat)
        zf = self._export()
        invoice = next(n for n in zf.namelist() if "_发票_1." in n)
        self.assertEqual(
            sorted(self._red_spans(zf.read(invoice))),
            sorted([(0, "01"), (0, f"订单号: {ALIPAY_ORDER_NO}"),
                    (0, f"交易单号: {wechat}")]),
        )

    def test_broken_file_kept_verbatim(self):
        item = self._item()
        self._attach(item, Attachment.KIND_INVOICE, "bad.pdf", b"not a pdf")
        self._attach(item, Attachment.KIND_PAYMENT, "p.png", self._png_bytes(),
                     order_no=ALIPAY_ORDER_NO)
        zf = self._export()
        invoice = next(n for n in zf.namelist() if "_发票_1." in n)
        self.assertEqual(zf.read(invoice), b"not a pdf")  # 标注失败原样放入
        payment = next(n for n in zf.namelist() if "_支付记录_1." in n)
        self.assertEqual(self._red_spans(zf.read(payment)), [(0, "01")])

    def test_item_without_invoice_stamps_payment_only(self):
        item = self._item()
        self._attach(item, Attachment.KIND_PAYMENT, "p.png", self._png_bytes(),
                     order_no=ALIPAY_ORDER_NO)
        zf = self._export()
        payment = next(n for n in zf.namelist() if "_支付记录_1." in n)
        self.assertEqual(self._red_spans(zf.read(payment)), [(0, "01")])

    def test_layout_matches_manual_spec(self):
        """布局对齐人工批次规格：hebo 22 pt 序号距底 12 pt。china-s 8.5 pt 单号距底 34 pt。"""
        item = self._item()
        self._attach(item, Attachment.KIND_INVOICE, "inv.pdf", self._pdf_bytes(1))
        self._attach(item, Attachment.KIND_PAYMENT, "p.png", self._png_bytes(300, 200),
                     order_no=ALIPAY_ORDER_NO)
        zf = self._export()
        invoice = next(n for n in zf.namelist() if "_发票_1." in n)
        doc = pymupdf.open(stream=zf.read(invoice), filetype="pdf")
        page = doc[0]
        height = page.rect.height
        spans = {
            span["text"]: span
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"] if span["color"] == 0xFF0000
        }
        seq = spans["01"]
        self.assertEqual(seq["font"], "Helvetica-Bold")
        self.assertEqual(seq["size"], 22.0)
        self.assertEqual(seq["origin"], (12.0, height - 12.0))
        order = spans[f"订单号: {ALIPAY_ORDER_NO}"]
        self.assertEqual(order["font"], "Heiti")
        self.assertEqual(order["size"], 8.5)
        self.assertEqual(order["origin"], (12.0, height - 34.0))

@override_settings(EXPECTED_INVOICE_TITLE="", EXPECTED_INVOICE_TAX_ID="")

@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class CheckItemTests(SubmissionTestCase):
    def _item(self, **kwargs):
        data = {
            "owner": self.student,
            "title": "书",
            "category": self.category,
            "batch": self.batch,
            "actual_amount": "100.00",
            "invoice_amount": "100.00",
        }
        data.update(kwargs)
        item = Item.objects.create(**data)
        item.refresh_from_db()  # 从库回读，字段值与真实请求路径一致（Decimal）
        return item

    def _payment(self, item, **kwargs):
        data = {"item": item, "kind": Attachment.KIND_PAYMENT,
                "file": "attachments/p.png", "amount": "100.00",
                "order_no": ALIPAY_ORDER_NO}
        data.update(kwargs)
        return Attachment.objects.create(**data)

    def test_missing_platform_order_no_warns(self):
        item = self._item()
        self._payment(item, order_no="")
        self.assertIn(
            "某笔支付的平台单号格式不符（支付宝订单号/微信交易单号/京东订单编号），报销时需抄到发票上",
            check_item(item),
        )

    def test_union_32_digit_order_no_passes(self):
        """银联等银行渠道单号 20 开头可到 32 位。"""
        item = self._item()
        self._payment(item, order_no="20" + "2" * 30)
        self.assertEqual(check_item(item), [])
        self._payment(item, order_no="20" + "1" * 31, amount=None)  # 33 位超出渠道上限
        self.assertIn("平台单号格式不符", check_item(item)[0])

    def test_wechat_28_digit_order_no_passes(self):
        """微信交易单号 4 开头 28 位。"""
        item = self._item()
        self._payment(item, order_no="4" + "5" * 27)
        self.assertEqual(check_item(item), [])
        self._payment(item, order_no="4" + "1" * 26)  # 27 位不足微信单号长度
        self.assertTrue(any("平台单号格式不符" in w for w in check_item(item)))

    def test_jd_order_no_passes(self):
        """京东订单编号 3 开头 12 或 16 位（生产数据实测两种长度）。"""
        item = self._item()
        self._payment(item, order_no="34" + "1" * 14)
        self.assertEqual(check_item(item), [])
        self._payment(item, order_no="34" + "1" * 10, amount=None)
        self.assertEqual(check_item(item), [])
        self._payment(item, order_no="3" + "1" * 13, amount=None)  # 14 位不是已知京东格式
        self.assertTrue(any("平台单号格式不符" in w for w in check_item(item)))

    def test_refund_net_amount_is_clean(self):
        item = self._item(actual_amount="179.00", invoice_amount="199.00")
        self._payment(item, amount="199.00")
        Attachment.objects.create(item=item, kind=Attachment.KIND_REFUND,
                                  file="attachments/r.png", amount="20.00")
        self.assertEqual(check_item(item), [])

    def test_attachment_totals_must_match_item_amounts(self):
        item = self._item(actual_amount="100.00", invoice_amount="100.00")
        self._payment(item, amount="90.00")
        Attachment.objects.create(item=item, kind=Attachment.KIND_INVOICE,
                                  file="attachments/i.png", amount="80.00")
        warnings = check_item(item)
        self.assertIn("支付与退款附件合计与填写的实付款不一致，请核对", warnings)
        self.assertIn("发票附件金额合计与填写的发票金额不一致，请核对", warnings)

    @override_settings(EXPECTED_INVOICE_TITLE="清华大学", EXPECTED_INVOICE_TAX_ID="12345678")
    def test_invoice_buyer_checked_against_expected(self):
        item = self._item()
        Attachment.objects.create(
            item=item, kind=Attachment.KIND_INVOICE, file="attachments/i.png",
            amount="100.00",
            ocr_data={"buyer_name": "北京大学", "buyer_id": "12345678"},
        )
        warnings = check_item(item)
        self.assertIn("发票抬头与要求不符", warnings)
        self.assertNotIn("发票税号与要求不符", warnings)

    @override_settings(EXPECTED_INVOICE_TITLE="清华大学", EXPECTED_INVOICE_TAX_ID="12345678")
    def test_missing_ocr_data_skips_buyer_checks(self):
        item = self._item()
        Attachment.objects.create(item=item, kind=Attachment.KIND_INVOICE,
                                  file="attachments/i.png", amount="100.00")
        self.assertNotIn("发票抬头与要求不符", check_item(item))
        self.assertNotIn("发票税号与要求不符", check_item(item))


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class AuditLogTests(SubmissionTestCase):
    def test_create_logs_full_snapshot(self):
        item = self._create_item(self.student)
        logs = AuditLog.objects.filter(action="create")
        self.assertEqual(logs.count(), 1)
        log = logs.get()
        self.assertEqual(log.actor, self.student)
        self.assertEqual(log.actor_name, "张三")
        self.assertEqual(log.item_pk, item.pk)
        self.assertEqual(log.snapshot["title"], "硬盘")
        self.assertEqual(log.snapshot["actual_amount"], "100.00")
        entries = log.snapshot["attachments"]
        self.assertEqual(len(entries), 2)
        self.assertEqual(
            sorted(entry["kind"] for entry in entries), ["invoice", "payment"]
        )
        self.assertTrue(entries[0]["file"])

    def test_update_logs_field_diff_only(self):
        item = self._create_item(self.student)
        before = AuditLog.objects.count()
        data = _item_data(
            category=self.category.pk,
            actual_amount="200.00",
            invoices=[_png_upload("invoice2.png")],
            payments=[_png_upload("payment2.png")],
        )
        self.client.post(f"/items/{item.pk}/edit/", data)
        item.refresh_from_db()
        self.assertEqual(AuditLog.objects.count(), before + 1)
        log = AuditLog.objects.filter(action="update").latest("id")
        self.assertEqual(log.actor, self.student)
        self.assertEqual(log.snapshot["actual_amount"], ["100.00", "200.00"])
        self.assertIn("新增 发票 invoice2.png", log.snapshot["attachments"])
        self.assertIn("新增 支付记录 payment2.png", log.snapshot["attachments"])

    def test_attachment_only_change_is_logged(self):
        item = self._create_item(self.student)
        attachment = item.attachments.get(kind=Attachment.KIND_PAYMENT)
        before = AuditLog.objects.count()
        self.client.force_login(self.student)
        self.client.post(
            f"/items/attachments/{attachment.pk}/edit/",
            {"kind": Attachment.KIND_PAYMENT, "amount": "55.50",
             "order_no": ALIPAY_ORDER_NO, "merchant_no": "", "invoice_no": ""},
        )
        self.assertEqual(AuditLog.objects.count(), before + 1)
        log = AuditLog.objects.filter(action="update").latest("id")
        self.assertNotIn("actual_amount", log.snapshot)
        self.assertIn("金额 由 空 改为 55.50", log.snapshot["attachments"])

    def test_no_op_save_writes_nothing(self):
        item = self._create_item(self.student)
        before = AuditLog.objects.count()
        item.save()
        item.save()
        self.assertEqual(AuditLog.objects.count(), before)

    def test_owner_delete(self):
        item = self._create_item(self.student)
        response = self.client.post(f"/items/{item.pk}/delete/")
        self.assertRedirects(response, "/")
        self.assertFalse(Item.objects.filter(pk=item.pk).exists())
        log = AuditLog.objects.filter(action="delete").latest("id")
        self.assertEqual(log.actor, self.student)
        self.assertEqual(log.item_pk, item.pk)
        self.assertIsNone(log.item)
        self.assertEqual(log.snapshot["title"], "硬盘")
        self.assertEqual(log.snapshot["status"], "pending")
        self.assertEqual(len(log.snapshot["attachments"]), 2)

    def test_staff_can_delete_any_item(self):
        item = self._create_item(self.student)
        self.client.force_login(self.staff)
        response = self.client.post(f"/items/{item.pk}/delete/")
        self.assertRedirects(response, "/")
        self.assertFalse(Item.objects.filter(pk=item.pk).exists())
        log = AuditLog.objects.filter(action="delete").latest("id")
        self.assertEqual(log.actor, self.staff)

    def test_delete_other_users_item_forbidden(self):
        item = self._create_item(self.student)
        self.client.force_login(self.student_b)
        response = self.client.post(f"/items/{item.pk}/delete/")
        self.assertEqual(response.status_code, 403)
        self.assertTrue(Item.objects.filter(pk=item.pk).exists())
        self.assertFalse(AuditLog.objects.filter(action="delete").exists())

    def _request_with_messages(self, user):
        request = self.factory.post("/")
        request.user = user
        request._messages = CookieStorage(request)
        return request

    def test_admin_bulk_action_logs_per_item(self):
        item1 = self._create_item(self.student)
        item2 = Item.objects.create(
            owner=self.student, title="书", category=self.category, batch=self.batch,
            actual_amount="50.00",
        )
        ma = admin.site._registry[Item]
        request = self._request_with_messages(self.staff)
        queryset = Item.objects.filter(pk__in=[item1.pk, item2.pk])
        ActorMiddleware(lambda r: ma.approve_items(r, queryset))(request)
        for item in (item1, item2):
            item.refresh_from_db()
            self.assertEqual(item.status, Item.STATUS_APPROVED)
        logs = AuditLog.objects.filter(action="update")
        self.assertEqual(logs.count(), 2)
        for log in logs:
            self.assertEqual(log.actor, self.staff)
            self.assertEqual(log.snapshot, {"status": ["pending", "approved"]})

    def test_shell_path_has_no_actor(self):
        item = Item.objects.create(
            owner=self.student, title="无请求", category=self.category, batch=self.batch,
            actual_amount="10.00",
        )
        log = AuditLog.objects.get(action="create", item_pk=item.pk)
        self.assertIsNone(log.actor)
        self.assertEqual(log.snapshot["attachments"], [])

    def test_update_page_shows_audit_section(self):
        item = self._create_item(self.student)
        response = self.client.get(f"/items/{item.pk}/edit/")
        self.assertContains(response, "操作记录")
        self.assertContains(response, "明细：硬盘")
        self.assertContains(response, "创建")

@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ItemDetailTests(SubmissionTestCase):
    """条目侧边栏片段：所有登录用户可见，编辑控件限本人/管理员。"""

    def test_panel_renders_images_and_groups(self):
        item = self._create_item(self.student)
        invoice = item.attachments.get(kind=Attachment.KIND_INVOICE)
        payment = item.attachments.get(kind=Attachment.KIND_PAYMENT)
        self.client.force_login(self.student)
        response = self.client.get(f"/items/{item.pk}/panel/")
        self.assertContains(response, f'src="/items/attachments/{invoice.pk}/file/"')
        self.assertContains(response, f'src="/items/attachments/{payment.pk}/file/"')
        self.assertContains(response, "发票</h4>")
        self.assertContains(response, "支付记录</h4>")
        self.assertNotContains(response, "退款记录</h4>")

    def test_panel_shows_refund_group_when_present(self):
        item = self._create_item(self.student, refunds=[_png_upload("r.png")])
        self.client.force_login(self.student)
        self.assertContains(self.client.get(f"/items/{item.pk}/panel/"), "退款记录</h4>")

    def test_panel_pdf_attachment_uses_lazy_iframe(self):
        item = self._create_item(self.student, invoices=[_png_upload("a.pdf")])
        pdf = item.attachments.get(kind=Attachment.KIND_INVOICE)
        self.client.force_login(self.student)
        response = self.client.get(f"/items/{item.pk}/panel/")
        self.assertContains(response, f'data-src="/items/attachments/{pdf.pk}/file/"')
        self.assertNotContains(response, f'<iframe src="/items/attachments/{pdf.pk}/file/"')

    def test_other_user_sees_readonly_panel(self):
        item = self._create_item(self.student)
        invoice = item.attachments.get(kind=Attachment.KIND_INVOICE)
        self.client.force_login(self.student_b)
        response = self.client.get(f"/items/{item.pk}/panel/")
        self.assertContains(response, f'src="/items/attachments/{invoice.pk}/file/"')
        self.assertContains(response, "操作记录")
        self.assertNotContains(response, 'id="panel-form"')
        self.assertNotContains(response, 'name="status"')

    def test_owner_and_staff_get_editable_panel(self):
        item = self._create_item(self.student)
        panel_url = f"/items/{item.pk}/panel/"
        self.client.force_login(self.student)
        response = self.client.get(panel_url)
        self.assertContains(response, 'id="panel-form"')
        self.assertNotContains(response, 'name="status"')
        self.assertNotContains(response, 'name="owner"')
        self.client.force_login(self.staff)
        response = self.client.get(panel_url)
        self.assertContains(response, 'name="status"')
        self.assertContains(response, 'name="owner"')

    def test_panel_requires_login(self):
        item = self._create_item(self.student)
        self.client.logout()
        self.assertEqual(self.client.get(f"/items/{item.pk}/panel/").status_code, 302)

    def test_board_requires_login(self):
        self.client.logout()
        self.assertEqual(self.client.get("/").status_code, 302)


class PairingTests(SimpleTestCase):
    """pair() 纯函数：配对规则 1-4 的可观察行为。"""

    @staticmethod
    def _record(rid, kind, amount=None, **extra):
        record = {"id": rid, "kind": kind, "amount": None if amount is None else Decimal(str(amount))}
        record.update(extra)
        return record

    def test_remark_matches_payment_order_no(self):
        result = pairing.pair([
            self._record(0, "invoice", "100", remark_order_no="X1"),
            self._record(1, "payment", "100", order_no="X1"),
        ])
        self.assertEqual(result["groups"], [[0, 1]])
        self.assertEqual(result["pending"], [])
        self.assertEqual(result["unmatched"], [])

    def test_remark_matches_merchant_no_and_merges_duplicates(self):
        result = pairing.pair([
            self._record(0, "invoice", "60", remark_order_no="M9"),
            self._record(1, "invoice", "40", remark_order_no="M9"),
            self._record(2, "payment", "60", merchant_no="M9"),
            self._record(3, "payment", "40", merchant_no="M9"),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2, 3]])

    def test_invoice_hitting_two_payment_groups_merges_them(self):
        result = pairing.pair([
            self._record(0, "invoice", "30", remark_order_no="X1"),
            self._record(1, "payment", "10", order_no="X1"),
            self._record(2, "payment", "20", merchant_no="X1"),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2]])

    def test_refund_attaches_by_order_no(self):
        result = pairing.pair([
            self._record(0, "invoice", "80", remark_order_no="X1"),
            self._record(1, "payment", "100", order_no="X1"),
            self._record(2, "refund", "20", order_no="X1"),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2]])
        self.assertEqual(result["unmatched"], [])

    def test_refund_attaches_by_amount_fallback(self):
        result = pairing.pair([
            self._record(0, "invoice", "80", remark_order_no="X1"),
            self._record(1, "payment", "100", order_no="X1"),
            self._record(2, "refund", "100"),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2]])

    def test_refund_without_payment_is_unmatched(self):
        result = pairing.pair([self._record(0, "refund", "5", order_no="NOPE")])
        self.assertEqual(result["unmatched"], [{"id": 0, "reason": "退款记录找不到对应支付记录"}])

    def test_amount_fallback_pairs_unique_candidates(self):
        result = pairing.pair([
            self._record(0, "invoice", "199"),
            self._record(1, "payment", "199"),
        ])
        self.assertEqual(result["groups"], [[0, 1]])

    def test_amount_fallback_multi_candidates_go_pending(self):
        result = pairing.pair([
            self._record(0, "invoice", "50"),
            self._record(1, "invoice", "50"),
            self._record(2, "payment", "50"),
            self._record(3, "payment", "50"),
        ])
        self.assertEqual(result["pending"], [[0, 1, 2, 3]])
        self.assertEqual(result["groups"], [])
        self.assertEqual(result["unmatched"], [])

    def test_amount_fallback_merges_combined_invoices(self):
        """合付：多张发票合计与一笔支付总额相等时并成一组（一笔支付对应多张发票）。"""
        result = pairing.pair([
            self._record(0, "invoice", "161.00"),
            self._record(1, "invoice", "254.94"),
            self._record(2, "invoice", "81.16"),
            self._record(3, "payment", "497.10"),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2, 3]])
        self.assertEqual(result["pending"], [])
        self.assertEqual(result["unmatched"], [])

    def test_combined_invoices_with_amount_gap_left_to_llm(self):
        """合付但发票合计与支付差额超容差：规则不并组，留给 LLM 兜底与人工确认。"""
        result = pairing.pair([
            self._record(0, "invoice", "161.00"),
            self._record(1, "invoice", "254.94"),
            self._record(2, "invoice", "81.16"),
            self._record(3, "payment", "496.96"),
        ])
        self.assertEqual(result["groups"], [])
        self.assertEqual(result["pending"], [])
        self.assertEqual(len(result["unmatched"]), 4)

    def test_subset_ambiguity_goes_pending(self):
        """组合多候选：单张 100 与 30+70 都配得上支付 100，整体待确认。"""
        result = pairing.pair([
            self._record(0, "invoice", "30"),
            self._record(1, "invoice", "70"),
            self._record(2, "invoice", "100"),
            self._record(3, "payment", "100"),
        ])
        self.assertEqual(result["groups"], [])
        self.assertEqual(sorted(result["pending"]), [[0, 1, 2, 3]])
        self.assertEqual(result["unmatched"], [])

    def test_invoice_shared_by_two_payments_goes_pending(self):
        result = pairing.pair([
            self._record(0, "invoice", "100"),
            self._record(1, "payment", "100"),
            self._record(2, "payment", "100"),
        ])
        self.assertEqual(result["groups"], [])
        self.assertEqual(sorted(result["pending"]), [[0, 1, 2]])

    def test_float_amounts_are_normalized(self):
        """金额混入识别原始 JSON 的 float：入口归一 Decimal，不炸并正常配平。"""
        result = pairing.pair([
            self._record(0, "invoice", 161.0),
            self._record(1, "invoice", 254.94),
            self._record(2, "invoice", 81.16),
            self._record(3, "payment", 497.1),
        ])
        self.assertEqual(result["groups"], [[0, 1, 2, 3]])

    def test_subset_enumeration_capped(self):
        """发票侧组件超上限：退化为单张比对，不枚举组合（16 张 1 元发票对 2 元支付不并组）。"""
        records = [self._record(i, "invoice", "1") for i in range(16)]
        records.append(self._record(99, "payment", "2"))
        result = pairing.pair(records)
        self.assertEqual(result["groups"], [])
        self.assertEqual(result["pending"], [])

    def test_support_record_goes_unmatched_with_hint(self):
        result = pairing.pair([
            self._record(0, "invoice", "100", remark_order_no="X1"),
            self._record(1, "payment", "100", order_no="X1"),
            self._record(2, "support", "1096.50"),
        ])
        self.assertEqual(result["groups"], [[0, 1]])
        self.assertEqual(result["unmatched"], [
            {"id": 2, "reason": "证明材料不参与自动配对，请加入对应组"},
        ])

    def test_all_unmatched_carry_reasons(self):
        result = pairing.pair([
            self._record(0, "invoice", "10"),
            self._record(1, "payment", "99"),
            self._record(2, "unknown"),
        ])
        self.assertEqual(result["unmatched"], [
            {"id": 0, "reason": "找不到对应支付记录"},
            {"id": 1, "reason": "找不到对应发票"},
            {"id": 2, "reason": "未能识别票据类型"},
        ])



@override_settings(DEEPSEEK_API_KEY="test-key")
class VisionDetectTests(SimpleTestCase):
    """detect / group_suggest / field_suggest 的解析与降级行为。"""

    @staticmethod
    def _response(content):
        return {"choices": [{"message": {"content": content}}]}

    def test_detect_returns_kind_and_fields(self):
        content = json.dumps({
            "kind": "payment", "amount": 12.5, "platform": "alipay",
            "order_no": ALIPAY_ORDER_NO, "merchant_no": "M123",
            "invoice_amount": None, "invoice_no": "", "remark_order_no": "",
            "buyer_name": "", "buyer_id": "", "items_summary": "",
        })
        with mock.patch("core.vision._post", return_value=self._response(content)) as post:
            result = vision.detect(PNG_1X1, "p.png")
        self.assertTrue(post.called)
        self.assertEqual(result["kind"], "payment")
        self.assertEqual(result["amount"], 12.5)
        self.assertEqual(result["platform"], "alipay")
        self.assertEqual(result["order_no"], ALIPAY_ORDER_NO)
        self.assertEqual(result["_source"], "image")

    def test_detect_invoice_passes_items_summary(self):
        content = json.dumps({
            "kind": "invoice", "invoice_amount": 100, "invoice_no": "INV1",
            "remark_order_no": "X1", "buyer_name": "清华大学", "buyer_id": "TAX1",
            "items_summary": "*信息技术服务*云存储", "amount": None,
            "platform": "unknown", "order_no": "", "merchant_no": "",
        })
        with mock.patch("core.vision._post", return_value=self._response(content)):
            result = vision.detect(PNG_1X1, "i.png")
        self.assertEqual(result["kind"], "invoice")
        self.assertEqual(result["invoice_amount"], 100)
        self.assertEqual(result["items_summary"], "*信息技术服务*云存储")
        self.assertEqual(result["remark_order_no"], "X1")

    def test_detect_garbage_returns_empty(self):
        with mock.patch("core.vision._post", return_value=self._response("not json")):
            self.assertEqual(vision.detect(PNG_1X1, "p.png"), {})

    def test_detect_unconfigured_returns_empty(self):
        with override_settings(DEEPSEEK_API_KEY=""):
            with mock.patch("core.vision._post") as post:
                self.assertEqual(vision.detect(PNG_1X1, "p.png"), {})
        self.assertFalse(post.called)

    def test_detect_routes_by_pdf_magic_without_extension(self):
        """腾讯云发票文件名超长被截断丢 .pdf 后缀、头部带脏字节：按 %PDF- 魔数走文本路径。"""
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((36, 60), "腾讯云发票 发票金额 99.00元 发票号码 26117000009900000001 " * 2)
        data = b"0okok" + doc.tobytes()  # 复刻线上脏头
        content = json.dumps({
            "kind": "invoice", "invoice_amount": 99.00, "invoice_no": "26117000009900000001",
            "remark_order_no": "", "buyer_name": "", "buyer_id": "",
            "items_summary": "腾讯云 轻量应用服务器", "amount": None,
            "platform": "unknown", "order_no": "", "merchant_no": "",
        })
        with mock.patch("core.vision._post", return_value=self._response(content)) as post:
            result = vision.detect(data, "999999999999_QCLOUD_qcloud_invoice_6ac3259edb9a917911740")
        self.assertTrue(post.called)
        self.assertEqual(result["kind"], "invoice")
        self.assertEqual(result["invoice_amount"], 99.00)
        self.assertEqual(result["_source"], "text")
        # 头部无 %PDF- 的图片仍走图片路径
        with mock.patch("core.vision._post", return_value=self._response(content)) as post:
            vision.detect(PNG_1X1, "photo")
        self.assertIn("image_url", json.dumps(post.call_args[0][0]["messages"][0]["content"]))

    def test_group_suggest_filters_unknown_ids_and_short_groups(self):
        content = json.dumps({"groups": [[7, 8, 99], ["bad"], [9]]})
        records = [
            {"id": 7, "kind": "invoice", "amount": "30", "order_no": "", "merchant_no": "", "remark_order_no": ""},
            {"id": 8, "kind": "payment", "amount": "30", "order_no": "Z", "merchant_no": "", "remark_order_no": ""},
            {"id": 9, "kind": "refund", "amount": "1", "order_no": "", "merchant_no": "", "remark_order_no": ""},
        ]
        with mock.patch("core.vision._post", return_value=self._response(content)):
            self.assertEqual(vision.group_suggest(records), [[7, 8]])

    def test_field_suggest_drops_invalid_category(self):
        content = json.dumps({"suggestions": [
            {"title": "打印费", "category_id": 999}, {"title": "书", "category_id": 1},
        ]})
        with mock.patch("core.vision._post", return_value=self._response(content)):
            result = vision.field_suggest(
                [[{"id": 0, "kind": "invoice", "amount": "10", "items_summary": "打印"}]],
                [{"id": 1, "name": "书籍", "description": ""}],
            )
        self.assertEqual(len(result["suggestions"]), 1)  # 超出组数的建议被截断
        self.assertEqual(result["suggestions"][0]["title"], "打印费")
        self.assertIsNone(result["suggestions"][0]["category_id"])  # 类别不在表中丢为 None

    def test_field_suggest_pads_short_response_to_group_count(self):
        content = json.dumps({"suggestions": [{"title": "打印费", "category_id": 1}]})
        groups = [[{"id": 0, "kind": "invoice", "amount": "10", "items_summary": "打印"}],
                  [{"id": 1, "kind": "payment", "amount": "10", "items_summary": ""}]]
        with mock.patch("core.vision._post", return_value=self._response(content)):
            result = vision.field_suggest(groups, [{"id": 1, "name": "书籍", "description": ""}])
        self.assertEqual(len(result["suggestions"]), len(groups))  # 模型少给时补 None 与组等长
        self.assertEqual(result["suggestions"][0], {"title": "打印费", "category_id": 1})
        self.assertIsNone(result["suggestions"][1])


class StagingSafeNameTests(SimpleTestCase):
    def test_long_name_keeps_extension(self):
        """超长文件名截断保留扩展名：丢后缀会让识别按文件名错走图片路径。"""
        long_pdf = "999999999999_QCLOUD_E02EA9TEST00000000000000_qcloud_invoice_6ac3259edb9a91791174" + "0123456789" * 5 + ".pdf"
        name = staging._safe_name(long_pdf)
        self.assertLessEqual(len(name), 80)
        self.assertTrue(name.endswith(".pdf"))

    def test_short_names_unchanged(self):
        self.assertEqual(staging._safe_name("发票 2026.pdf"), "发票_2026.pdf")
        self.assertEqual(staging._safe_name(""), "file")
        no_suffix = staging._safe_name("q" * 100)
        self.assertEqual(len(no_suffix), 80)
        self.assertFalse(no_suffix.endswith(".pdf"))



@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BatchStageTests(SubmissionTestCase):
    def test_stage_returns_session_and_persists_files(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/stage/",
            {"files": [_png_upload("发票 1.png"), _png_upload("payment.pdf")]},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertRegex(data["session_id"], r"^[0-9a-f]{32}$")
        self.assertEqual([f["name"] for f in data["files"]], ["发票_1.png", "payment.pdf"])
        from django.conf import settings
        session_dir = Path(settings.MEDIA_ROOT) / "tmp" / "batch" / data["session_id"]
        self.assertEqual(len(list(session_dir.iterdir())), 3)  # meta.json + 2 文件

    def test_stage_rejects_bad_extension_and_login(self):
        self.client.force_login(self.student)
        response = self.client.post("/items/batch/stage/", {"files": [_png_upload("a.gif")]})
        self.assertEqual(response.status_code, 400)
        self.client.logout()
        self.assertEqual(self.client.post("/items/batch/stage/", {}).status_code, 302)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BatchSubmitTests(SubmissionTestCase):

    def _post_batch(self, session_id, groups):
        return self.client.post(
            "/items/batch/submit/",
            json.dumps({"session_id": session_id, "groups": groups}),
            content_type="application/json",
        )

    def test_submit_creates_pending_items_with_attachments(self):
        session_id, file_ids = self._stage([
            _png_upload("invoice.png"), _png_upload("payment.png"), _png_upload("refund.png"),
            _png_upload("invoice2.png"), _png_upload("payment2.png"),
        ])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [
            {
                "title": "硬盘", "category": self.category.pk,
                "actual_amount": "80.00", "invoice_amount": "100.00",
                "files": [
                    {"file_id": file_ids[0], "kind": "invoice", "amount": "100.00",
                     "invoice_no": "INV1", "ocr": {"buyer_name": "清华大学"}},
                    {"file_id": file_ids[1], "kind": "payment", "amount": "100.00",
                     "order_no": ALIPAY_ORDER_NO},
                    {"file_id": file_ids[2], "kind": "refund", "amount": "20.00",
                     "merchant_no": "R1"},
                ],
            },
            {
                "title": "打印费", "category": self.category.pk,
                "actual_amount": "5.00", "invoice_amount": "5.00",
                "files": [
                    {"file_id": file_ids[3], "kind": "invoice", "amount": "5.00"},
                    {"file_id": file_ids[4], "kind": "payment", "amount": "5.00"},
                ],
            },
        ])
        self.assertRedirects(response, "/")
        self.assertEqual(Item.objects.count(), 2)
        item = Item.objects.order_by("id").first()
        self.assertEqual(item.owner, self.student)
        self.assertEqual(item.title, "硬盘")
        self.assertEqual(item.status, Item.STATUS_PENDING)
        self.assertEqual(str(item.actual_amount), "80.00")
        self.assertEqual(item.attachments.count(), 3)
        payment = item.attachments.get(kind=Attachment.KIND_PAYMENT)
        self.assertEqual(str(payment.amount), "100.00")
        self.assertEqual(payment.order_no, ALIPAY_ORDER_NO)
        invoice = item.attachments.get(kind=Attachment.KIND_INVOICE)
        self.assertEqual(invoice.ocr_data, {"buyer_name": "清华大学"})
        log = AuditLog.objects.get(action="create", item_pk=item.pk)
        self.assertEqual(len(log.snapshot["attachments"]), 3)

    def test_submit_support_file_attached_without_amount_effects(self):
        """证明材料随组提交：附件落库但不参与金额（腾讯云充值流水场景）。"""
        session_id, file_ids = self._stage([
            _png_upload("invoice.png"), _png_upload("payment.png"), _png_upload("flow.png"),
        ])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [
            {
                "title": "腾讯云服务器", "category": self.category.pk,
                "actual_amount": "656.13", "invoice_amount": "656.13",
                "files": [
                    {"file_id": file_ids[0], "kind": "invoice", "amount": "656.13",
                     "invoice_no": "INV-QC"},
                    {"file_id": file_ids[1], "kind": "payment", "amount": "1096.50",
                     "order_no": ALIPAY_ORDER_NO},
                    {"file_id": file_ids[2], "kind": "support", "amount": "1096.50",
                     "ocr": {"items_summary": "腾讯云充值 1096.50 元入余额"}},
                ],
            },
        ])
        self.assertRedirects(response, "/")
        item = Item.objects.get(title="腾讯云服务器")
        kinds = {a.kind: a.amount for a in item.attachments.all()}
        self.assertEqual(kinds["support"], Decimal("1096.50"))
        # 证明材料不参与净额核对：支付 1096.50 对实付 656.13 只有一条金额警告
        warnings = [w for w in check_item(item) if "实付款不一致" in w]
        self.assertEqual(len(warnings), 1)

    def test_session_cleared_after_submit(self):
        session_id, file_ids = self._stage(
            [_png_upload("invoice.png"), _png_upload("payment.png")])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [{
            "title": "打印费", "category": self.category.pk,
            "actual_amount": "5.00", "invoice_amount": "5.00",
            "files": [
                {"file_id": file_ids[0], "kind": "invoice", "amount": "5.00"},
                {"file_id": file_ids[1], "kind": "payment", "amount": "5.00"},
            ],
        }])
        self.assertRedirects(response, "/")
        from django.conf import settings
        self.assertFalse(
            (Path(settings.MEDIA_ROOT) / "tmp" / "batch" / session_id).exists())

    def test_solo_invoice_group_submits_with_zero_actual(self):
        session_id, file_ids = self._stage([_png_upload("invoice.png")])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [{
            "title": "只有发票", "category": self.category.pk,
            "actual_amount": "0.00", "invoice_amount": "10.00",
            "files": [{"file_id": file_ids[0], "kind": "invoice", "amount": "10.00"}],
        }])
        self.assertRedirects(response, "/")
        item = Item.objects.latest("id")
        self.assertEqual(str(item.actual_amount), "0.00")
        self.assertEqual(item.attachments.count(), 1)

    def test_payment_only_group_rejected_without_items(self):
        session_id, file_ids = self._stage([_png_upload("payment.png")])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [{
            "title": "只有支付", "category": self.category.pk,
            "actual_amount": "10.00",
            "files": [{"file_id": file_ids[0], "kind": "payment", "amount": "10.00"}],
        }])
        self.assertEqual(response.status_code, 400)
        self.assertIn("缺少发票", response.json()["error"])
        self.assertFalse(Item.objects.exists())

    def test_submit_rejects_foreign_session(self):
        session_id, _ = self._stage([_png_upload("invoice.png")])
        other = User.objects.create_user(username="other24", password="x")
        self.client.force_login(other)
        response = self._post_batch(session_id, [{
            "title": "盗用", "category": self.category.pk, "actual_amount": "1.00",
            "files": [{"file_id": "0" * 32, "kind": "invoice", "amount": "1.00"}],
        }])
        self.assertEqual(response.status_code, 400)
        self.assertIn("不属于当前用户", response.json()["error"])

    def test_full_refund_group_accepts_zero_actual(self):
        # 全退款组净额 0：actual_amount 为 JSON 数字 0 时不能被当缺失拒掉
        session_id, file_ids = self._stage([
            _png_upload("invoice.png"), _png_upload("payment.png"), _png_upload("refund.png")])
        self.client.force_login(self.student)
        response = self._post_batch(session_id, [{
            "title": "全退款", "category": self.category.pk,
            "actual_amount": 0, "invoice_amount": "20.00",
            "files": [
                {"file_id": file_ids[0], "kind": "invoice", "amount": "20.00"},
                {"file_id": file_ids[1], "kind": "payment", "amount": "20.00",
                 "order_no": ALIPAY_ORDER_NO},
                {"file_id": file_ids[2], "kind": "refund", "amount": "20.00",
                 "order_no": ALIPAY_ORDER_NO},
            ],
        }])
        self.assertRedirects(response, "/")
        item = Item.objects.latest("id")
        self.assertEqual(str(item.actual_amount), "0.00")

    def test_amount_beyond_model_limits_rejected(self):
        session_id, file_ids = self._stage(
            [_png_upload("invoice.png"), _png_upload("payment.png")])
        self.client.force_login(self.student)
        base = {
            "title": "越界", "category": self.category.pk,
            "files": [
                {"file_id": file_ids[0], "kind": "invoice", "amount": "10.00"},
                {"file_id": file_ids[1], "kind": "payment", "amount": "10.00"},
            ],
        }
        for bad in ("100000000", "1.005"):
            with self.subTest(amount=bad):
                response = self._post_batch(
                    session_id, [{**base, "actual_amount": bad}])
                self.assertEqual(response.status_code, 400)
                self.assertIn("最多两位小数", response.json()["error"])
        self.assertFalse(Item.objects.exists())

    def test_requires_login(self):
        response = self.client.post("/items/batch/submit/", "{}", content_type="application/json")
        self.assertEqual(response.status_code, 302)


@override_settings(DEEPSEEK_API_KEY="test-key", MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BatchPairViewTests(SubmissionTestCase):

    def _post_pair(self, session_id, records):
        return self.client.post(
            "/items/batch/pair/",
            json.dumps({"session_id": session_id, "records": records}),
            content_type="application/json",
        )

    def test_rule_groups_returned_when_unconfigured(self):
        session_id, _ = self._stage([_png_upload("a.png"), _png_upload("b.png")])
        self.client.force_login(self.student)
        with override_settings(DEEPSEEK_API_KEY=""):
            response = self._post_pair(session_id, [
                {"id": 0, "kind": "invoice", "amount": "100", "remark_order_no": "X1"},
                {"id": 1, "kind": "payment", "amount": "100", "order_no": "X1"},
            ])
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["groups"], [[0, 1]])
        self.assertEqual(data["pending"], [])
        self.assertEqual(data["unmatched"], [])
        self.assertEqual(data["suggestions"], {})

    def test_unmatched_invoices_extracted_to_solo_groups(self):
        session_id, _ = self._stage([_png_upload("a.png"), _png_upload("b.png")])
        self.client.force_login(self.student)
        with override_settings(DEEPSEEK_API_KEY="", MEDIA_ROOT=TEMP_MEDIA_ROOT):
            response = self._post_pair(session_id, [
                {"id": 0, "kind": "invoice", "amount": "100", "invoice_no": "INV1"},
                {"id": 1, "kind": "invoice", "amount": "100", "invoice_no": "INV1"},
                {"id": 2, "kind": "payment", "amount": "999", "order_no": "Z9"},
            ])
        data = response.json()
        self.assertEqual(data["groups"], [])
        self.assertEqual(sorted(data["solo"]), [[0, 1]])  # 同发票号码重复拍摄并为一组
        self.assertEqual([entry["id"] for entry in data["unmatched"]], [2])

    def test_llm_groups_merged_with_suggestions(self):
        session_id, _ = self._stage([_png_upload("a.png"), _png_upload("b.png")])
        self.client.force_login(self.student)
        with mock.patch("core.vision.group_suggest", return_value=[[0, 1]]) as group_mock:
            with mock.patch(
                "core.vision.field_suggest",
                return_value={"suggestions": [{"title": "打印费", "category_id": self.category.pk}]},
            ) as field_mock:
                response = self._post_pair(session_id, [
                    {"id": 0, "kind": "invoice", "amount": "30"},
                    {"id": 1, "kind": "payment", "amount": "35", "order_no": "Z9"},
                ])
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["groups"], [[0, 1]])
        self.assertEqual(data["unmatched"], [])
        self.assertEqual(group_mock.call_args[0][0][0]["id"], 0)
        self.assertEqual(field_mock.call_args[0][0][0][0]["id"], 0)
        self.assertEqual(data["suggestions"]["suggestions"][0]["title"], "打印费")

    def test_llm_group_without_payment_dropped(self):
        """LLM 给出的纯发票组无法提交：丢弃，发票走单边发票组流程。"""
        session_id, _ = self._stage([_png_upload("a.png"), _png_upload("b.png"), _png_upload("c.png")])
        self.client.force_login(self.student)
        with mock.patch("core.vision.group_suggest", return_value=[[0, 1]]):
            response = self._post_pair(session_id, [
                {"id": 0, "kind": "invoice", "amount": "30"},
                {"id": 1, "kind": "invoice", "amount": "35"},
                {"id": 2, "kind": "payment", "amount": "70"},
            ])
        data = response.json()
        self.assertEqual(data["groups"], [])
        self.assertEqual(sorted(data["solo"]), [[0], [1]])
        self.assertEqual([entry["id"] for entry in data["unmatched"]], [2])

    def test_requires_login(self):
        response = self.client.post("/items/batch/pair/", "{}", content_type="application/json")
        self.assertEqual(response.status_code, 302)


@override_settings(DEEPSEEK_API_KEY="test-key")
class VisionRequestShapeTests(SimpleTestCase):
    def test_group_suggest_sends_content_as_part_list(self):
        """group_suggest/field_suggest 的 content 参数为结构化消息体。"""
        from core import vision
        response = {"choices": [{"message": {"content": '{"groups": []}'}}]}
        with mock.patch.object(vision, "_post", return_value=response) as post_mock:
            vision.group_suggest([{"id": 0, "kind": "invoice", "amount": "10"}])
        payload = post_mock.call_args[0][0]
        content = payload["messages"][0]["content"]
        self.assertIsInstance(content, list)
        self.assertEqual(content[0]["type"], "text")

    def test_merge_page_results_sums_payments_and_keeps_first_no(self):
        """多笔支付合一份 PDF：金额相加，单号取首个。类型不一致退回首页。"""
        from core import vision
        pages = [
            {"kind": "payment", "amount": 100.0, "order_no": "A1"},
            {"kind": "payment", "amount": 50.0, "order_no": "B2"},
            {"kind": "payment", "amount": 10.0, "order_no": "C3"},
        ]
        merged = vision._merge_page_results(pages)
        self.assertEqual(merged["kind"], "payment")
        self.assertEqual(merged["amount"], 160.0)
        self.assertEqual(merged["order_no"], "A1")
        mixed = [{"kind": "invoice", "amount": 20.0}, {"kind": "payment", "amount": 20.0}]
        self.assertEqual(vision._merge_page_results(mixed)["kind"], "invoice")

    def test_merge_page_results_invoice_takes_first_amount(self):
        """发票多页（重复打印）金额不累加。"""
        from core import vision
        pages = [
            {"kind": "invoice", "invoice_amount": 160.0},
            {"kind": "invoice", "invoice_amount": 160.0},
        ]
        self.assertEqual(vision._merge_page_results(pages)["invoice_amount"], 160.0)


@override_settings(DEEPSEEK_API_KEY="test-key", MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BatchAgentRoundTests(SubmissionTestCase):
    def _stage_two(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/stage/",
            {"files": [_png_upload("发票.png"), _png_upload("支付.png")]},
        )
        data = response.json()
        return data["session_id"], data["files"][0]["id"], data["files"][1]["id"]

    def test_reread_corrections_merge_and_regroup(self):
        session_id, invoice_id, payment_id = self._stage_two()
        corrected = {"kind": "invoice", "invoice_amount": 100.0, "remark_order_no": "X1",
                     "handwritten_notes": "30 划掉改 100"}
        unchanged = {"kind": "payment", "amount": 100.0, "order_no": "X1"}
        with mock.patch("core.vision.group_suggest", return_value=[]), \
             mock.patch("core.vision.reread", side_effect=[corrected, unchanged]) as reread_mock:
            response = self.client.post(
                "/items/batch/agent-round/",
                json.dumps({"session_id": session_id, "records": [
                    {"id": 0, "file_id": invoice_id, "kind": "invoice", "amount": "30",
                     "filename": "发票.png"},
                    {"id": 1, "file_id": payment_id, "kind": "payment", "amount": "100",
                     "order_no": "X1", "filename": "支付.png"},
                ]}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["groups"], [[0, 1]])
        self.assertEqual(data["done"], True)
        self.assertEqual(data["records"][0]["amount"], "100.0")
        self.assertEqual(data["records"][0]["remark_order_no"], "X1")
        self.assertTrue(data["actions"])
        self.assertEqual(reread_mock.call_count, 2)  # 初始两条都未决，各重读一次
        self.assertEqual(reread_mock.call_args_list[0][0][1], "发票.png")

    def test_reread_amount_correction_does_not_crash_repair(self):
        """重读修正支付金额为 float 时，第二轮配对正常运行。"""
        session_id, invoice_id, payment_id = self._stage_two()
        invoice_reread = {"kind": "invoice", "invoice_amount": 30.0}
        payment_reread = {"kind": "payment", "amount": 30.0, "order_no": "Z9"}
        with mock.patch("core.vision.group_suggest", return_value=[]), \
             mock.patch("core.vision.reread", side_effect=[invoice_reread, payment_reread]):
            response = self.client.post(
                "/items/batch/agent-round/",
                json.dumps({"session_id": session_id, "records": [
                    {"id": 0, "file_id": invoice_id, "kind": "invoice", "amount": "30",
                     "filename": "发票.png"},
                    {"id": 1, "file_id": payment_id, "kind": "payment", "amount": "999",
                     "order_no": "Z9", "filename": "支付.png"},
                ]}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["groups"], [[0, 1]])  # 修正后发票 30 = 支付 30，金额兜底成组
        self.assertEqual(data["records"][1]["amount"], "30.0")

    def test_no_improvement_ends_with_done(self):
        session_id, invoice_id, payment_id = self._stage_two()
        with mock.patch("core.vision.group_suggest", return_value=[]), \
             mock.patch("core.vision.reread", return_value={}):
            response = self.client.post(
                "/items/batch/agent-round/",
                json.dumps({"session_id": session_id, "records": [
                    {"id": 0, "file_id": invoice_id, "kind": "invoice", "amount": "30",
                     "filename": "发票.png"},
                    {"id": 1, "file_id": payment_id, "kind": "payment", "amount": "999",
                     "order_no": "Z9", "filename": "支付.png"},
                ]}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["groups"], [])
        self.assertEqual(data["done"], True)  # 本轮无改善，提前止损
        self.assertIn("重读失败", data["actions"][0]["text"])

    def test_requires_login(self):
        response = self.client.post(
            "/items/batch/agent-round/", "{}", content_type="application/json")
        self.assertEqual(response.status_code, 302)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BoardTableTests(SubmissionTestCase):
    def test_defaults_to_latest_batch_with_seq(self):
        older = Batch.objects.create(name="旧批次")
        Batch.objects.create(name="新批次")
        Item.objects.create(owner=self.student, title="旧条目", category=self.category,
                            batch=older, actual_amount="1.00", position=1)
        item = self._create_item(self.student)  # 归最新批（新批次）
        item.status = Item.STATUS_APPROVED
        item.save()
        self.client.force_login(self.student)
        response = self.client.get("/")
        self.assertContains(response, item.title)
        self.assertNotContains(response, "旧条目")
        self.assertContains(response, '<td class="num seq">1</td>')

    def test_seq_matches_export_approved_only(self):
        """混合状态批次：approved 占连续序号且与导出一致，待审/退回不占号。"""
        approved = self._create_item(self.student, title="已通过")
        approved.status = Item.STATUS_APPROVED
        approved.save()
        self._create_item(self.student, title="待审核")
        rejected = self._create_item(self.student, title="已退回")
        rejected.status = Item.STATUS_REJECTED
        rejected.save()
        self.client.force_login(self.staff)
        response = self.client.get("/")
        self.assertContains(response, '<td class="num seq">1</td>')
        self.assertContains(response, '<td class="num seq">—</td>')
        self.assertNotContains(response, '<td class="num seq">2</td>')
        # 同一批次导出：序号连续、只含 approved，与表格序号严格一致
        sheet = load_workbook(BytesIO(self.client.get(
            "/export/excel/", {"batch": approved.batch.pk}).content)).active
        self.assertEqual(sheet.cell(row=2, column=1).value, 1)
        self.assertEqual(sheet.cell(row=2, column=2).value, "已通过")
        self.assertNotIn("待审核", [cell.value for row in sheet.iter_rows() for cell in row])

    def test_batch_switch_and_filters(self):
        other_category = Category.objects.create(name="其他类别", order=5)
        item_a = self._create_item(self.student, title="硬盘")
        item_b = self._create_item(self.staff, title="会议费", category=other_category.pk)
        item_b.status = Item.STATUS_APPROVED
        item_b.save()
        self.client.force_login(self.student)
        self.assertContains(self.client.get("/"), "硬盘")
        self.assertContains(self.client.get("/"), "会议费")
        response = self.client.get("/", {"status": "approved"})
        self.assertNotContains(response, "硬盘")
        self.assertContains(response, "会议费")
        response = self.client.get("/", {"category": other_category.pk})
        self.assertNotContains(response, "硬盘")
        response = self.client.get("/", {"payer": self.staff.pk})
        self.assertNotContains(response, "硬盘")
        response = self.client.get("/", {"q": "硬盘"})
        self.assertNotContains(response, "会议费")

    def test_staff_sees_admin_controls_student_does_not(self):
        self._create_item(self.student)
        self.client.force_login(self.staff)
        response = self.client.get("/")
        self.assertContains(response, "新建批次")
        self.assertContains(response, "导出 Excel")
        self.assertContains(response, "导出 zip")
        self.client.force_login(self.student)
        response = self.client.get("/")
        self.assertNotContains(response, "新建批次")
        self.assertNotContains(response, "导出 Excel")

    def test_export_buttons_follow_download_origin(self):
        """配置 DOWNLOAD_ORIGIN 后导出按钮指向下载域。未配置保持相对链接。"""
        self.client.force_login(self.staff)
        with override_settings(DOWNLOAD_ORIGIN="https://dl.example.com"):
            html = self.client.get("/").content.decode()
        self.assertIn(f'href="https://dl.example.com/export/excel/?batch={self.batch.pk}"', html)
        self.assertIn(f'href="https://dl.example.com/export/zip/?batch={self.batch.pk}"', html)
        html = self.client.get("/").content.decode()
        self.assertIn(f'href="/export/zip/?batch={self.batch.pk}"', html)

    def test_no_batches_empty_state(self):
        Batch.objects.all().delete()
        self.client.force_login(self.staff)
        self.assertContains(self.client.get("/"), "创建批次")
        self.client.force_login(self.student)
        self.assertContains(self.client.get("/"), "请联系管理员创建批次")


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ItemFieldUpdateTests(SubmissionTestCase):
    def _post_field(self, item, field, value, as_user=None):
        self.client.force_login(as_user or self.student)
        return self.client.post(
            f"/items/{item.pk}/field/",
            json.dumps({"field": field, "value": value}),
            content_type="application/json",
        )

    def test_owner_updates_title(self):
        item = self._create_item(self.student)
        response = self._post_field(item, "title", "新明细")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        item.refresh_from_db()
        self.assertEqual(item.title, "新明细")

    def test_other_student_forbidden(self):
        item = self._create_item(self.student)
        self.assertEqual(self._post_field(item, "title", "x", as_user=self.student_b).status_code, 403)

    def test_student_cannot_change_status_or_owner(self):
        item = self._create_item(self.student)
        self.assertEqual(self._post_field(item, "status", "approved").status_code, 403)
        self.assertEqual(self._post_field(item, "owner", self.student_b.pk).status_code, 403)
        item.refresh_from_db()
        self.assertEqual(item.status, Item.STATUS_PENDING)
        self.assertEqual(item.owner, self.student)

    def test_staff_changes_status_and_owner_with_audit(self):
        item = self._create_item(self.student)
        response = self._post_field(item, "status", "approved", as_user=self.staff)
        self.assertEqual(response.status_code, 200)
        response = self._post_field(item, "owner", self.student_b.pk, as_user=self.staff)
        self.assertEqual(response.status_code, 200)
        item.refresh_from_db()
        self.assertEqual(item.status, Item.STATUS_APPROVED)
        self.assertEqual(item.owner, self.student_b)
        log = AuditLog.objects.filter(action="update").latest("id")
        self.assertEqual(log.snapshot, {"owner": ["张三", "李四"]})

    def test_category_change_appends_to_end(self):
        other_category = Category.objects.create(name="其他类别", order=5)
        first = self._create_item(self.student)   # position 1
        second = self._create_item(self.student)  # position 2
        response = self._post_field(first, "category", other_category.pk)
        self.assertEqual(response.status_code, 200)
        first.refresh_from_db()
        self.assertEqual(first.category_id, other_category.pk)
        self.assertEqual(first.position, 3)  # next_position()：落到目标类别末尾

    def test_invalid_amount_rejected(self):
        item = self._create_item(self.student)
        for bad in ("abc", "1.005", "100000000"):
            with self.subTest(amount=bad):
                response = self._post_field(item, "actual_amount", bad)
                self.assertEqual(response.status_code, 400)
        item.refresh_from_db()
        self.assertEqual(str(item.actual_amount), "100.00")


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ItemReorderTests(SubmissionTestCase):
    def _reorder(self, order, category, as_user=None):
        self.client.force_login(as_user or self.staff)
        return self.client.post(
            "/items/reorder/",
            json.dumps({"category": category.pk, "order": order}),
            content_type="application/json",
        )

    def test_staff_reorder_updates_positions_and_seq(self):
        item_a = self._create_item(self.student)   # position 1
        item_b = self._create_item(self.student)   # position 2
        item_c = self._create_item(self.staff)     # position 3
        response = self._reorder([item_c.pk, item_a.pk, item_b.pk], self.category)
        self.assertEqual(response.status_code, 200)
        for item, position in ((item_c, 4), (item_a, 5), (item_b, 6)):
            item.refresh_from_db()
            self.assertEqual(item.position, position)
        self.client.force_login(self.staff)
        content = self.client.get("/").content.decode()
        pks = re.findall(r'data-pk="(\d+)"', content)
        self.assertEqual(pks, [str(item_c.pk), str(item_a.pk), str(item_b.pk)])

    def test_student_forbidden(self):
        item = self._create_item(self.student)
        self.assertEqual(self._reorder([item.pk], self.category, as_user=self.student_b).status_code, 403)

    def test_order_must_cover_exactly(self):
        item_a = self._create_item(self.student)
        item_b = self._create_item(self.student)
        other_category = Category.objects.create(name="其他类别", order=5)
        other = Item.objects.create(owner=self.student, title="他类", category=other_category,
                                    batch=self.batch, actual_amount="1.00", position=10)
        self.assertEqual(self._reorder([item_a.pk], self.category).status_code, 400)  # 缺行
        self.assertEqual(self._reorder([item_a.pk, item_b.pk, other.pk], self.category).status_code, 400)  # 跨类别
        self.assertEqual(self._reorder([item_a.pk, item_b.pk, item_b.pk], self.category).status_code, 400)  # 重复
        self.assertEqual(self._reorder([item_a.pk, 999999], self.category).status_code, 400)  # 不存在


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ExportBatchTests(SubmissionTestCase):
    def _approve(self, item):
        item.status = Item.STATUS_APPROVED
        item.save(update_fields=["status"])
        return item

    def test_batches_isolated_with_batch_filename(self):
        item = self._approve(self._create_item(self.student))
        other_batch = Batch.objects.create(name="第二批")
        self._approve(Item.objects.create(
            owner=self.student, title="他批条目", category=self.category,
            batch=other_batch, actual_amount="9.00", position=99,
        ))
        self.client.force_login(self.staff)
        response = self.client.get("/export/excel/", {"batch": self.batch.pk})
        self.assertIn("报销汇总_2025 报销批次.xlsx", _content_disposition(response))
        sheet = load_workbook(BytesIO(response.content)).active
        titles = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
        self.assertIn("硬盘", titles)
        self.assertNotIn("他批条目", titles)
        # 缺省 batch 取最新批（第二批创建最晚）
        response = self.client.get("/export/excel/")
        sheet = load_workbook(BytesIO(response.content)).active
        titles = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
        self.assertIn("他批条目", titles)
        self.assertNotIn("硬盘", titles)

    def test_zip_filename_contains_batch_name(self):
        Item.objects.create(owner=self.student, title="后提交排前", category=self.category,
                            batch=self.batch, actual_amount="2.00", position=1,
                            status=Item.STATUS_APPROVED)
        Item.objects.create(owner=self.student, title="先提交排后", category=self.category,
                            batch=self.batch, actual_amount="1.00", position=2,
                            status=Item.STATUS_APPROVED)
        self.client.force_login(self.staff)
        response = self.client.get("/export/zip/", {"batch": self.batch.pk})
        self.assertIn("报销材料_2025 报销批次.zip", _content_disposition(response))

    def test_excel_seq_follows_position(self):
        Item.objects.create(owner=self.student, title="后提交排前", category=self.category,
                            batch=self.batch, actual_amount="2.00", position=1,
                            status=Item.STATUS_APPROVED)
        Item.objects.create(owner=self.student, title="先提交排后", category=self.category,
                            batch=self.batch, actual_amount="1.00", position=2,
                            status=Item.STATUS_APPROVED)
        self.client.force_login(self.staff)
        sheet = load_workbook(BytesIO(self.client.get(
            "/export/excel/", {"batch": self.batch.pk}).content)).active
        self.assertEqual(sheet.cell(row=2, column=1).value, 1)
        self.assertEqual(sheet.cell(row=2, column=2).value, "后提交排前")
        self.assertEqual(sheet.cell(row=3, column=1).value, 2)
        self.assertEqual(sheet.cell(row=3, column=2).value, "先提交排后")

    def test_export_without_batches_yields_empty_files(self):
        Batch.objects.all().delete()
        self.client.force_login(self.staff)
        response = self.client.get("/export/excel/")
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 1)  # 仅表头
        response = self.client.get("/export/zip/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(zipfile.ZipFile(BytesIO(response.content)).namelist()), 0)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ArchivedBatchLockTests(SubmissionTestCase):
    """存档批次只读：条目/附件/调序写操作对 owner 与 staff 一致拒绝，界面不渲染编辑控件。"""

    def setUp(self):
        super().setUp()
        self.archived = Batch.objects.create(name="2025 存档批", archived=True)
        self.a_item = Item.objects.create(
            owner=self.student, batch=self.archived, category=self.category,
            title="存档条目", actual_amount=Decimal("10"), position=Item.next_position(),
        )

    def test_field_update_rejected_for_owner_and_staff(self):
        for user in (self.student, self.staff):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                response = self.client.post(
                    f"/items/{self.a_item.pk}/field/",
                    data=json.dumps({"field": "title", "value": "改存档"}),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 403)
        self.a_item.refresh_from_db()
        self.assertEqual(self.a_item.title, "存档条目")

    def test_update_delete_reorder_rejected_for_staff(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            f"/items/{self.a_item.pk}/edit/",
            _item_data(category=self.category.pk, owner=self.staff.pk, status=Item.STATUS_PENDING),
        )
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.post(f"/items/{self.a_item.pk}/delete/").status_code, 403)
        response = self.client.post(
            "/items/reorder/",
            data=json.dumps({"order": [self.a_item.pk], "category": self.category.pk}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)
        self.assertTrue(Item.objects.filter(pk=self.a_item.pk).exists())

    def test_attachment_write_rejected_for_staff(self):
        att = Attachment.objects.create(
            item=self.a_item, kind=Attachment.KIND_INVOICE,
            file=SimpleUploadedFile("f.png", b"x", content_type="image/png"),
            amount=Decimal("10"),
        )
        self.client.force_login(self.staff)
        self.assertEqual(
            self.client.post(f"/items/attachments/{att.pk}/edit/", {"amount": "11"}).status_code, 403
        )
        self.assertEqual(
            self.client.post(f"/items/attachments/{att.pk}/delete/").status_code, 403
        )
        self.assertTrue(Attachment.objects.filter(pk=att.pk).exists())

    def test_board_and_panel_render_readonly_for_staff(self):
        self.client.force_login(self.staff)
        response = self.client.get("/", {"batch": self.archived.pk})
        self.assertContains(response, "（已存档）")
        self.assertNotContains(response, '<td class="drag-handle"')
        self.assertNotContains(response, 'class="danger small row-del"')
        panel = self.client.get(f"/items/{self.a_item.pk}/panel/")
        self.assertIsNone(panel.context.get("form"))
        self.assertContains(panel, "已存档")

    def test_create_rejected_when_latest_batch_archived(self):
        Batch.objects.all().update(archived=True)
        self.client.force_login(self.student)
        response = self.client.post("/items/new/", _item_data(category=self.category.pk))
        self.assertRedirects(response, "/")
        self.assertEqual(Item.objects.count(), 1)  # 仅剩 setUp 的存档条目，未新建


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class BatchAddTests(SubmissionTestCase):
    def test_staff_creates_batch(self):
        self.client.force_login(self.staff)
        response = self.client.post("/batches/add/", {"name": "20261008 报销"})
        self.assertRedirects(response, "/")
        self.assertTrue(Batch.objects.filter(name="20261008 报销").exists())

    def test_student_forbidden(self):
        self.client.force_login(self.student)
        self.assertEqual(self.client.post("/batches/add/", {"name": "x"}).status_code, 403)
        self.assertFalse(Batch.objects.filter(name="x").exists())

    def test_duplicate_name_rejected(self):
        self.client.force_login(self.staff)
        response = self.client.post("/batches/add/", {"name": self.batch.name}, follow=True)
        self.assertContains(response, "已存在")
        self.assertEqual(Batch.objects.filter(name=self.batch.name).count(), 1)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT, DEEPSEEK_API_KEY="test-key")
class ImportBatchCommandTests(SubmissionTestCase):
    def setUp(self):
        super().setUp()
        self.root = User.objects.create_user("root", password="pw", is_superuser=True)

    def call(self, directory, *args):
        out = StringIO()
        call_command("import_batch", str(directory), *args, stdout=out)
        return out.getvalue()

    def _write(self, tmp, person, files):
        directory = Path(tmp) / person if person else Path(tmp)
        directory.mkdir(parents=True, exist_ok=True)
        paths = []
        for name in files:
            path = directory / name
            path.write_bytes(PNG_1X1)
            paths.append(path)
        return paths

    def test_person_dir_creates_pending_item(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, "张三", ["invoice.jpg", "payment.jpg"])
            with mock.patch("core.vision.detect") as detect:
                detect.side_effect = [
                    {"kind": "invoice", "invoice_amount": 100, "invoice_no": "INV1",
                     "items_summary": "移动硬盘"},
                    {"kind": "payment", "amount": 100, "order_no": ALIPAY_ORDER_NO},
                ]
                with mock.patch(
                    "core.vision.field_suggest",
                    return_value={"suggestions": [{"title": "硬盘", "category_id": self.category.pk}]},
                ):
                    out = self.call(tmp, "--batch", "导入测试批")
        batch = Batch.objects.get(name="导入测试批")
        item = Item.objects.get()
        self.assertEqual(item.owner, self.student)  # 人名张三按 first_name 匹配
        self.assertEqual(item.status, Item.STATUS_PENDING)
        self.assertEqual(item.batch, batch)
        self.assertEqual(item.title, "硬盘")
        self.assertEqual(str(item.actual_amount), "100.00")
        self.assertEqual(item.attachments.count(), 2)
        self.assertIn("建条 1 条", out)

    def test_unknown_person_falls_back_to_executor(self):
        root = self.root
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, "王五", ["invoice.jpg", "payment.jpg"])
            with mock.patch("core.vision.detect") as detect:
                detect.side_effect = [
                    {"kind": "invoice", "invoice_amount": 50},
                    {"kind": "payment", "amount": 50, "order_no": ALIPAY_ORDER_NO},
                ]
                with mock.patch("core.vision.field_suggest", return_value={"suggestions": []}):
                    out = self.call(tmp, "--batch", "导入测试批")
        item = Item.objects.get()
        self.assertEqual(item.owner, root)
        self.assertIn("人名 王五 无匹配账号", out)
        self.assertIn("归执行者 root", out)

    def test_single_sided_file_reported_not_created(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, None, ["01 论文录用通知.pdf"])
            with mock.patch("core.vision.detect", return_value={"kind": "invoice", "invoice_amount": 200}):
                out = self.call(tmp, "--batch", "导入测试批")
        self.assertEqual(Item.objects.count(), 0)
        self.assertIn("找不到对应支付记录", out)
        self.assertIn("未成条目", out)


    def test_negative_amount_file_reported_not_created(self):
        """识别出负数金额的文件按识别失败报告，导入不中断、不建条目。"""
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, "张三", ["invoice.jpg", "payment.jpg"])
            with mock.patch("core.vision.detect") as detect:
                detect.side_effect = [
                    {"kind": "invoice", "invoice_amount": -5},
                    {"kind": "payment", "amount": 100},
                ]
                out = self.call(tmp, "--batch", "导入测试批")
        self.assertEqual(Item.objects.count(), 0)
        self.assertIn("识别金额为负", out)
        self.assertIn("找不到对应发票", out)

    def test_missing_amount_attachments_do_not_break_import(self):
        """附件识别不到金额时导入不崩溃，无法配对的按未成条目报告。"""
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, "张三", ["invoice.jpg", "payment.jpg"])
            with mock.patch("core.vision.detect") as detect:
                detect.side_effect = [
                    {"kind": "invoice", "invoice_amount": 100},
                    {"kind": "payment"},
                ]
                out = self.call(tmp, "--batch", "导入测试批")
        self.assertEqual(Item.objects.count(), 0)
        self.assertIn("payment.jpg", out)

    def test_dry_run_creates_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._write(tmp, "张三", ["invoice.jpg", "payment.jpg"])
            with mock.patch("core.vision.detect") as detect:
                detect.side_effect = [
                    {"kind": "invoice", "invoice_amount": 100},
                    {"kind": "payment", "amount": 100, "order_no": ALIPAY_ORDER_NO},
                ]
                with mock.patch("core.vision.field_suggest", return_value={"suggestions": []}):
                    out = self.call(tmp, "--batch", "导入测试批", "--dry-run")
        self.assertEqual(Item.objects.count(), 0)
        self.assertFalse(Batch.objects.filter(name="导入测试批").exists())
        self.assertIn("DRY-RUN 未落库", out)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class ExportFormulaSafetyTests(SubmissionTestCase):
    """导出 Excel 的用户文本一律写成文本单元格，防公式注入。"""

    def test_title_starting_with_equals_is_not_a_formula(self):
        item = Item.objects.create(
            owner=self.student, title="=HYPERLINK(\"http://evil\",\"点我\")",
            category=self.category, batch=self.batch,
            actual_amount="1.00", position=1, status=Item.STATUS_APPROVED,
        )
        self.client.force_login(self.staff)
        response = self.client.get("/export/excel/", {"batch": self.batch.pk})
        sheet = load_workbook(BytesIO(response.content)).active
        cell = sheet.cell(row=2, column=2)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(str(cell.value).startswith("="))
        self.assertIsNotNone(item)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class RobustInputTests(SubmissionTestCase):
    """行内接口对非法输入返回 400，不抛 500。"""

    def test_field_update_non_dict_json_is_400(self):
        item = self._create_item(self.student)
        self.client.force_login(self.student)
        response = self.client.post(
            f"/items/{item.pk}/field/", "[]", content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_field_update_bad_category_value_is_400(self):
        item = self._create_item(self.student)
        self.client.force_login(self.student)
        for bad in ("abc", "[1]", "0", "-3"):
            response = self.client.post(
                f"/items/{item.pk}/field/",
                json.dumps({"field": "category", "value": bad}),
                content_type="application/json")
            self.assertEqual(response.status_code, 400, bad)
            self.assertIn("类别无效", response.json()["error"])

    def test_field_update_bad_owner_value_is_400(self):
        item = self._create_item(self.student)
        self.client.force_login(self.staff)
        response = self.client.post(
            f"/items/{item.pk}/field/",
            json.dumps({"field": "owner", "value": "abc"}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("付款人无效", response.json()["error"])

    def test_reorder_bad_category_value_is_400(self):
        item = self._create_item(self.student)
        self.client.force_login(self.staff)
        response = self.client.post(
            "/items/reorder/",
            json.dumps({"category": "abc", "order": [item.pk]}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_reorder_non_dict_json_is_400(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            "/items/reorder/", "null", content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_prefill_kind_error_lists_all_kinds(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/prefill/", {"file": _png_upload("a.png"), "kind": "bad"})
        self.assertEqual(response.status_code, 400)
        message = response.json()["error"]
        for kind in ("invoice", "payment", "refund", "support"):
            self.assertIn(kind, message)


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class DraftCheckTests(SubmissionTestCase):
    """规则本体走服务端核对接口（core/validation.py）。"""

    def test_item_draft_check_returns_server_warnings(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/draft-check/",
            json.dumps({
                "actual_amount": "5.00", "invoice_amount": "10.00",
                "attachments": [
                    {"kind": "payment", "amount": "5.00", "order_no": "bad"},
                    {"kind": "invoice", "amount": "10.00"},
                ],
            }),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        warnings = response.json()["warnings"]
        self.assertTrue(any("平台单号格式不符" in w for w in warnings))

    def test_batch_group_check_reports_missing_fields(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/check/",
            json.dumps({
                "groups": [
                    {"title": "", "category": "", "actual_amount": "",
                     "files": [{"kind": "invoice", "amount": "10.00"}]},
                    {"title": "有支付", "category": self.category.pk,
                     "actual_amount": "10.00",
                     "files": [{"kind": "payment", "amount": "10.00"}]},
                ],
            }),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        results = response.json()["groups"]
        self.assertIn("缺少明细", results[0]["issues"])
        self.assertIn("类别无效或未选择", results[0]["issues"])
        self.assertEqual(results[0]["state"], "solo")
        self.assertIn("缺少发票", results[1]["issues"])
        self.assertEqual(results[1]["state"], "no_invoice")

    def test_batch_group_check_mismatch_state(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/check/",
            json.dumps({"groups": [{
                "title": "金额不平", "category": self.category.pk,
                "actual_amount": "80.00",
                "files": [{"kind": "invoice", "amount": "100.00"},
                          {"kind": "payment", "amount": "80.00"}],
            }]}),
            content_type="application/json")
        self.assertEqual(response.json()["groups"][0]["state"], "mismatch")
        self.assertEqual(response.json()["groups"][0]["issues"], [])

    def test_draft_check_skips_file_only_rules(self):
        """未保存草稿无文件，深色模式等文件类检查自动跳过。"""
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/draft-check/",
            json.dumps({
                "actual_amount": "10.00", "invoice_amount": "10.00",
                "attachments": [
                    {"kind": "invoice", "amount": "10.00"},
                    {"kind": "payment", "amount": "10.00", "order_no": ALIPAY_ORDER_NO},
                ],
            }),
            content_type="application/json")
        self.assertEqual(response.json()["warnings"], [])


@override_settings(DEEPSEEK_API_KEY="test-key")
class MergePageResultsTests(SimpleTestCase):
    """多页识别合并：金额取值、类型判定、support 有意义判定。"""

    def test_merge_takes_first_nonempty_invoice_amount(self):
        from core import vision
        pages = [
            {"kind": "invoice", "invoice_amount": None, "invoice_no": "INV1"},
            {"kind": "invoice", "invoice_amount": 160.0, "invoice_no": ""},
        ]
        merged = vision._merge_page_results(pages, "invoice")
        self.assertEqual(merged["invoice_amount"], 160.0)
        self.assertEqual(merged["invoice_no"], "INV1")

    def test_merge_uses_caller_kind_for_prefill_path(self):
        """prefill 路径的识别结果不含 kind 字段，仍按调用方指定类型合并金额。"""
        from core import vision
        pages = [
            {"amount": 100.0, "order_no": "A1"},
            {"amount": 50.0, "order_no": ""},
        ]
        merged = vision._merge_page_results(pages, "payment")
        self.assertEqual(merged["amount"], 150.0)
        self.assertEqual(merged["order_no"], "A1")

    def test_support_result_with_only_items_summary_is_kept(self):
        from core import vision
        result = {"amount": None, "items_summary": "充值 100 元入余额"}
        self.assertTrue(vision._meaningful("support", result))
        self.assertFalse(vision._meaningful("support", {"amount": None, "items_summary": ""}))


@override_settings(MEDIA_ROOT=TEMP_MEDIA_ROOT)
class NonFiniteAmountTests(SubmissionTestCase):
    """金额为 NaN、Infinity 等非有限值时不得 500：宽松解析当缺金额，提交报 400。"""

    def test_draft_check_non_finite_amount_is_not_500(self):
        self.client.force_login(self.student)
        for bad in ("sNaN", "NaN", "Infinity"):
            response = self.client.post(
                "/items/draft-check/",
                json.dumps({
                    "actual_amount": "10.00", "invoice_amount": "10.00",
                    "attachments": [{"kind": "payment", "amount": bad}],
                }),
                content_type="application/json")
            self.assertEqual(response.status_code, 200, bad)

    def test_batch_group_check_non_finite_amount_is_not_500(self):
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/check/",
            json.dumps({"groups": [{
                "title": "非有限金额", "category": self.category.pk,
                "actual_amount": "sNaN",
                "files": [{"kind": "invoice", "amount": "10.00"},
                          {"kind": "payment", "amount": "NaN"}],
            }]}),
            content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["groups"][0]["state"], "ok")

    def test_submit_non_finite_amount_is_400(self):
        session_id, file_ids = self._stage(
            [_png_upload("invoice.png"), _png_upload("payment.png")])
        self.client.force_login(self.student)
        response = self.client.post(
            "/items/batch/submit/",
            json.dumps({"session_id": session_id, "groups": [{
                "title": "非有限金额", "category": self.category.pk,
            "actual_amount": "10.00", "invoice_amount": "10.00",
            "files": [
                {"file_id": file_ids[0], "kind": "invoice", "amount": "10.00"},
                {"file_id": file_ids[1], "kind": "payment", "amount": "sNaN"},
            ]}]}),
            content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("金额格式不正确", response.json()["error"])
        self.assertFalse(Item.objects.exists())
