import json
import zipfile
from decimal import Decimal
from io import BytesIO
from itertools import groupby
from types import SimpleNamespace

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files import File
from django.db import transaction
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from openpyxl import Workbook
from openpyxl.styles import Font

from . import pairing, staging, stamping, vision
from .attachments import KIND_LABELS, attachment_groups, build_attachment, decimal_or_none, plan_decimal
from .audit import attachment_entries, attachment_entry, mark_attachments
from .forms import AttachmentForm, ItemPanelForm
from .models import Attachment, Batch, Category, Item
from .suggest import category_payloads, llm_merge_groups, pair_summary, suggest_member
from .validation import check_item
from .vision import prefill as vision_prefill, validate_upload

STATUS_FILTERS = {Item.STATUS_PENDING, Item.STATUS_APPROVED, Item.STATUS_REJECTED}
AMOUNT_FORMAT = "0.00_ "

# 上传文件字段名、POST 字段前缀、附件类型 三元组
UPLOAD_FIELDS = (
    ("invoices", "new_invoice", Attachment.KIND_INVOICE),
    ("payments", "new_payment", Attachment.KIND_PAYMENT),
    ("refunds", "new_refund", Attachment.KIND_REFUND),
    ("supports", "new_support", Attachment.KIND_SUPPORT),
)


def _is_fetch(request):
    return request.headers.get("x-requested-with") == "fetch"


def _first_error(form, errors):
    for error in errors:
        return error
    for field_errors in form.errors.values():
        for error in field_errors:
            return str(error)
    return "保存失败"


@login_required
def board(request):
    """表格主视图：按批次展示全部条目，序号=该批 approved 条目的导出序号；筛选/行内编辑/拖拽调序由前端承载。"""
    batches = Batch.objects.all()
    current = None
    batch_id = request.GET.get("batch", "")
    if batch_id.isdigit():
        current = batches.filter(pk=batch_id).first()
    if current is None:
        current = batches.first()

    rows, payers, seen_payers = [], [], set()
    if current is not None:
        items = (
            current.items.select_related("owner", "category")
            .prefetch_related("attachments")
            .order_by("category__order", "category__id", "position", "id")
        )
        seq = 0
        for item in items:
            counts = {"invoice": 0, "payment": 0, "refund": 0, "support": 0}
            for attachment in item.attachments.all():
                if attachment.kind in counts:
                    counts[attachment.kind] += 1
            if item.owner_id not in seen_payers:
                seen_payers.add(item.owner_id)
                payers.append(item.owner)
            if item.status == Item.STATUS_APPROVED:
                seq += 1  # 与导出一致：只对 approved 连续编号
                row_seq = seq
            else:
                row_seq = None  # 待审/退回不占号，展示为 —
            rows.append({
                "seq": row_seq,
                "item": item,
                "can_edit": (request.user.is_staff or item.owner_id == request.user.id) and not current.archived,
                "counts": counts,
            })

    status = request.GET.get("status", "")
    category_id = request.GET.get("category", "")
    payer_id = request.GET.get("payer", "")
    query = request.GET.get("q", "").strip()
    if status in STATUS_FILTERS:
        rows = [row for row in rows if row["item"].status == status]
    if category_id.isdigit():
        rows = [row for row in rows if str(row["item"].category_id) == category_id]
    if payer_id.isdigit():
        rows = [row for row in rows if str(row["item"].owner_id) == payer_id]
    if query:
        rows = [row for row in rows if query in row["item"].title]

    return render(request, "core/board.html", {
        "batches": batches,
        "current_batch": current,
        "categories": Category.objects.all(),
        "payers": payers,
        "rows": rows,
        "status": status,
        "category_id": category_id,
        "payer_id": payer_id,
        "q": query,
        "status_choices": Item.STATUS_CHOICES,
        "download_origin": settings.DOWNLOAD_ORIGIN,
    })


@login_required
def item_create(request):
    batch = Batch.objects.first()
    if request.method == "POST":
        if batch is None:
            messages.error(request, "请先创建批次")
            return redirect("board")
        if batch.archived:
            messages.error(request, "最新批次已存档，请新建批次后再提交")
            return redirect("board")
        form = ItemPanelForm(request.POST, request.FILES, staff=request.user.is_staff)
        attachments, errors = _collect_new_attachments(request)
        if form.is_valid() and not errors:
            item = form.save(commit=False)
            if not request.user.is_staff:
                item.owner = request.user
            item.batch = batch
            item.position = Item.next_position()
            mark_attachments(item, [attachment_entry(a) for a in attachments], [])
            item.save()
            for attachment in attachments:
                attachment.item = item
                attachment.save()
            if _is_fetch(request):
                return JsonResponse({"ok": True, "warnings": check_item(item)})
            for warning in check_item(item):
                messages.warning(request, warning)
            messages.success(request, "已提交，等待管理员审核。")
            return redirect("board")
        for error in errors:
            form.add_error(None, error)
        if _is_fetch(request):
            return JsonResponse({"error": _first_error(form, errors)}, status=400)
    else:
        form = ItemPanelForm(staff=request.user.is_staff)
    return render(request, "core/item_panel.html", {
        "form": form,
        "is_create": True,
        "can_edit": not (batch and batch.archived),
        "is_staff": request.user.is_staff,
    })


@login_required
def item_update(request, pk):
    item = get_object_or_404(Item, pk=pk)
    if item.owner_id != request.user.id and not request.user.is_staff:
        raise PermissionDenied("只能修改自己的条目")
    _reject_archived(item)
    if request.method == "POST":
        form = ItemPanelForm(request.POST, request.FILES, instance=item, staff=request.user.is_staff)
        attachments, errors = _collect_new_attachments(request)
        if form.is_valid() and not errors:
            before = attachment_entries(item)
            item = form.save(commit=False)
            if not request.user.is_staff:
                item.owner = request.user
            mark_attachments(item, before + [attachment_entry(a) for a in attachments], before)
            item.save()
            for attachment in attachments:
                attachment.item = item
                attachment.save()
            if _is_fetch(request):
                return JsonResponse({"ok": True, "warnings": check_item(item)})
            for warning in check_item(item):
                messages.warning(request, warning)
            messages.success(request, "已更新。")
            return redirect("board")
        for error in errors:
            form.add_error(None, error)
        if _is_fetch(request):
            return JsonResponse({"error": _first_error(form, errors)}, status=400)
    else:
        form = ItemPanelForm(instance=item, staff=request.user.is_staff)
    return render(request, "core/item_panel.html", {
        "form": form,
        "item": item,
        "is_create": False,
        "can_edit": not item.batch.archived,
        "is_staff": request.user.is_staff,
        "audit_logs": item.audit_logs.all(),
        "attachment_groups": attachment_groups(item),
        "warnings": check_item(item),
    })


@login_required
def item_panel(request, pk):
    """条目侧边栏片段：登录即可看（含附件与留痕），编辑控件限本人/staff。"""
    item = get_object_or_404(
        Item.objects.select_related("owner", "category", "batch"), pk=pk
    )
    can_edit = (request.user.is_staff or item.owner_id == request.user.id) and not item.batch.archived
    context = {
        "item": item,
        "is_create": False,
        "can_edit": can_edit,
        "is_staff": request.user.is_staff,
        "audit_logs": item.audit_logs.all(),
        "attachment_groups": attachment_groups(item),
        "warnings": check_item(item),
    }
    if can_edit:
        context["form"] = ItemPanelForm(instance=item, staff=request.user.is_staff)
    return render(request, "core/item_panel.html", context)


@login_required
@require_POST
def item_field_update(request, pk):
    """表格行内单字段保存；状态/付款人仅 staff。"""
    item = get_object_or_404(Item, pk=pk)
    if item.owner_id != request.user.id and not request.user.is_staff:
        raise PermissionDenied("只能修改自己的条目")
    _reject_archived(item)
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    field = payload.get("field") if isinstance(payload, dict) else None
    value = payload.get("value") if isinstance(payload, dict) else None

    student_fields = {"title", "category", "actual_amount", "invoice_amount"}
    staff_fields = {"status", "owner"}
    if field in staff_fields and not request.user.is_staff:
        raise PermissionDenied("只有管理员可以修改审核状态或付款人")
    if field not in student_fields | staff_fields:
        return JsonResponse({"error": "不允许修改的字段"}, status=400)

    try:
        if field == "title":
            item.title = str(value or "").strip()[:200]
        elif field == "category":
            category = Category.objects.filter(pk=value).first()
            if category is None:
                return JsonResponse({"error": "类别无效"}, status=400)
            item.category = category
            item.position = Item.next_position()  # 行内改类别：落到目标类别末尾
        elif field in ("actual_amount", "invoice_amount"):
            try:
                setattr(item, field, plan_decimal(value, "金额"))
            except ValueError as exc:
                return JsonResponse({"error": str(exc)}, status=400)
        elif field == "status":
            if value not in STATUS_FILTERS:
                return JsonResponse({"error": "状态无效"}, status=400)
            item.status = value
        elif field == "owner":
            owner = User.objects.filter(pk=value, is_active=True).first()
            if owner is None:
                return JsonResponse({"error": "付款人无效"}, status=400)
            item.owner = owner
        item.full_clean()
    except ValidationError as exc:
        message = exc.message_dict and next(iter(exc.message_dict.values()))[0] or exc.messages[0]
        return JsonResponse({"error": message}, status=400)
    item.save()
    return JsonResponse({"ok": True, "warnings": check_item(item)})


@login_required
@require_POST
def item_reorder(request):
    """staff 拖拽调序：order 必须恰为该类别下当前批次的全部条目。"""
    if not request.user.is_staff:
        raise PermissionDenied("只有管理员可以调整顺序")
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    order = payload.get("order") if isinstance(payload, dict) else None
    category = Category.objects.filter(pk=payload.get("category")).first() if isinstance(payload, dict) else None
    if category is None or not isinstance(order, list) or not order:
        return JsonResponse({"error": "参数无效"}, status=400)
    pks = []
    for pk in order:
        if not isinstance(pk, int) or isinstance(pk, bool):
            return JsonResponse({"error": "order 元素必须是条目 id"}, status=400)
        pks.append(pk)
    if len(set(pks)) != len(pks):
        return JsonResponse({"error": "order 存在重复条目"}, status=400)
    items = {obj.pk: obj for obj in Item.objects.filter(pk__in=pks)}
    if len(items) != len(pks):
        return JsonResponse({"error": "order 含不存在的条目"}, status=400)
    if any(obj.batch.archived for obj in items.values()):
        return JsonResponse({"error": "已存档批次不可调整顺序"}, status=403)
    if any(obj.category_id != category.pk for obj in items.values()):
        return JsonResponse({"error": "order 含跨类别的条目"}, status=400)
    batch_id = next(iter(items.values())).batch_id
    expected = set(
        Item.objects.filter(batch_id=batch_id, category=category).values_list("pk", flat=True)
    )
    if set(pks) != expected:
        return JsonResponse({"error": "order 必须覆盖该类别下的全部条目"}, status=400)
    with transaction.atomic():
        position = Item.next_position()
        for pk in pks:
            item = items[pk]
            item.position = position
            item.save()
            position += 1
    return JsonResponse({"ok": True})


@login_required
@require_POST
def item_delete(request, pk):
    item = get_object_or_404(Item, pk=pk)
    if item.owner_id != request.user.id and not request.user.is_staff:
        raise PermissionDenied("只能删除自己的条目")
    _reject_archived(item)
    mark_attachments(item, attachment_entries(item), [])
    item.delete()
    if _is_fetch(request):
        return JsonResponse({"ok": True})
    messages.success(request, "已删除。")
    return redirect("board")


@login_required
@require_POST
def batch_add(request):
    if not request.user.is_staff:
        raise PermissionDenied("只有管理员可以创建批次")
    name = (request.POST.get("name") or "").strip()[:100]
    if not name:
        messages.error(request, "批次名称不能为空")
        return redirect("board")
    if Batch.objects.filter(name=name).exists():
        messages.error(request, f"批次 {name} 已存在")
        return redirect("board")
    Batch.objects.create(name=name)
    messages.success(request, f"已创建批次 {name}")
    return redirect("board")


@login_required
def file_serve(request, pk):
    attachment = get_object_or_404(Attachment, pk=pk)
    return FileResponse(attachment.file.open("rb"), as_attachment=False)


@login_required
@require_POST
def attachment_update(request, pk):
    attachment = _owned_attachment(request, pk)
    item = attachment.item
    before = attachment_entries(item)
    form = AttachmentForm(request.POST, instance=attachment)
    if form.is_valid():
        form.save()
        mark_attachments(item, attachment_entries(item), before)
        item.save()
        if _is_fetch(request):
            return JsonResponse({"ok": True})
        messages.success(request, "附件已更新。")
    else:
        error = next(iter(form.errors.values()))[0]
        if _is_fetch(request):
            return JsonResponse({"error": f"附件保存失败：{error}"}, status=400)
        messages.error(request, f"附件保存失败：{error}")
    return redirect("board")


@login_required
@require_POST
def attachment_delete(request, pk):
    attachment = _owned_attachment(request, pk)
    item = attachment.item
    before = attachment_entries(item)
    attachment.file.delete(save=False)
    attachment.delete()
    mark_attachments(item, attachment_entries(item), before)
    item.save()
    if _is_fetch(request):
        return JsonResponse({"ok": True})
    messages.success(request, "已删除附件。")
    return redirect("board")


@require_POST
@login_required
def prefill(request):
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "缺少 file 字段"}, status=400)
    kind = request.POST.get("kind", "")
    if kind not in KIND_LABELS:
        return JsonResponse({"error": "kind 必须是 invoice/payment/refund"}, status=400)
    error = validate_upload(upload)
    if error:
        return JsonResponse({"error": error}, status=400)
    return JsonResponse(vision_prefill(upload.read(), upload.name, kind))


@login_required
def batch_create(request):
    """批量配对提交页：混合上传 → 识别 → 配对 → 确认创建。"""
    categories = Category.objects.all()
    notes = [f"{category.name}：{category.description}" for category in categories if category.description]
    return render(request, "core/batch_form.html", {
        "vision_configured": vision.configured(),
        "categories": category_payloads(),
        "category_notes": "；".join(notes),
    })


@require_POST
@login_required
def batch_stage(request):
    """批量页暂存上传：文件落配对会话目录，返回 session_id 与文件清单（ADR-0008）。"""
    uploads = [f for f in request.FILES.getlist("files") if f.name]
    if not uploads:
        return JsonResponse({"error": "缺少 files"}, status=400)
    if len(uploads) > staging.MAX_SESSION_FILES:
        return JsonResponse({"error": f"一次最多上传 {staging.MAX_SESSION_FILES} 个文件"}, status=400)
    for upload in uploads:
        error = validate_upload(upload)
        if error:
            return JsonResponse({"error": f"{upload.name}：{error}"}, status=400)
    session_id, files = staging.create_session(request.user.id, uploads)
    return JsonResponse({"session_id": session_id, "files": files})


def _session_and_file(request, payload):
    """从请求体解析配对会话与文件路径；非法抛 ValueError（文案可直接进 400）。"""
    session = staging.session_dir(payload.get("session_id"), request.user.id)
    return session, staging.file_path(session, payload.get("file_id"))


@require_POST
@login_required
def batch_detect(request):
    """批量页逐张识别（读配对会话暂存）：判类型并提取字段，不落库。"""
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    try:
        _, path = _session_and_file(request, payload)
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    return JsonResponse(vision.detect(path.read_bytes(), staging.file_label(path)))


def _parse_records(raw_records):
    """解析前端提交的记录列表为配对输入；结构非法抛 ValueError（文案）。"""
    if not isinstance(raw_records, list):
        raise ValueError("缺少 records 列表")
    records = []
    for raw in raw_records:
        if not isinstance(raw, dict):
            raise ValueError("records 元素必须是对象")
        records.append({
            "id": raw.get("id"),
            "file_id": raw.get("file_id") or "",
            "kind": raw.get("kind"),
            "amount": decimal_or_none(raw.get("amount")),
            "order_no": raw.get("order_no") or "",
            "merchant_no": raw.get("merchant_no") or "",
            "invoice_no": raw.get("invoice_no") or "",
            "remark_order_no": raw.get("remark_order_no") or "",
            "handwritten_notes": raw.get("handwritten_notes") or "",
            "filename": raw.get("filename") or "",
            "items_summary": raw.get("items_summary") or "",
        })
    return records


def _json_records(records):
    """records 深拷贝为 JSON 安全结构（Decimal 金额转字符串）。"""
    return [{
        **record,
        "amount": str(record["amount"]) if record.get("amount") is not None else None,
    } for record in records]


def _solo_groups(unmatched, by_id):
    """未配对发票按发票号码聚成单边发票组（同号重复拍摄同组）；返回 (solo, 剩余 unmatched)。"""
    solos, rest = {}, []
    for entry in unmatched:
        record = by_id.get(entry["id"])
        if record is not None and record.get("kind") == Attachment.KIND_INVOICE:
            solos.setdefault(record.get("invoice_no") or f"__{entry['id']}", []).append(entry["id"])
        else:
            rest.append(entry)
    return list(solos.values()), rest


def _pair_records(records):
    """规则配对 + LLM 兜底 + 单边发票组提取。

    返回 (groups, pending, solo, unmatched, by_id)。
    """
    result = pairing.pair(records)
    groups, pending, unmatched = result["groups"], result["pending"], result["unmatched"]
    by_id = {record["id"]: record for record in records}
    if vision.configured():
        groups, pending, unmatched = llm_merge_groups(by_id, groups, pending, unmatched)
    solo, unmatched = _solo_groups(unmatched, by_id)
    return groups, pending, solo, unmatched, by_id


def _field_suggestions(groups, by_id):
    """成组记录的明细/类别建议；未配置或无组返回空 dict。"""
    if vision.configured() and groups:
        return vision.field_suggest(
            [[suggest_member(by_id[rid]) for rid in group] for group in groups], category_payloads()
        )
    return {}


@require_POST
@login_required
def batch_pair(request):
    """规则配对 + LLM 兜底 + 单边发票组提取：返回分组、待确认、单边发票组、未配对与字段建议。"""
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    try:
        staging.session_dir(payload.get("session_id"), request.user.id)
        records = _parse_records(payload.get("records"))
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    groups, pending, solo, unmatched, by_id = _pair_records(records)
    return JsonResponse({
        "groups": groups,
        "pending": pending,
        "solo": solo,
        "unmatched": unmatched,
        "suggestions": _field_suggestions(groups, by_id),
    })


AGENT_MAX_ROUNDS = 3
AGENT_MAX_REREADS = 6
_AGENT_UNRESOLVED_KINDS = (Attachment.KIND_INVOICE, Attachment.KIND_PAYMENT)
_AGENT_FIELDS = ("kind", "amount", "order_no", "merchant_no", "invoice_no",
                 "remark_order_no", "handwritten_notes")


def _unresolved_ids(pending, solo, unmatched, by_id):
    """待确认组、单边发票组与未配对里的发票/支付记录 id（智能体的重读对象，去重保序）。"""
    ids = [rid for group in pending for rid in group]
    ids += [rid for group in solo for rid in group]
    ids += [entry["id"] for entry in unmatched if by_id[entry["id"]].get("kind") in _AGENT_UNRESOLVED_KINDS]
    seen, result = set(), []
    for rid in ids:
        if rid not in seen:
            seen.add(rid)
            result.append(rid)
    return result


@require_POST
@login_required
def batch_agent_round(request):
    """配对智能体一轮（ADR-0008）：定向重读未决文件，合并修正字段后重跑配对。

    请求 {session_id, records}；返回 {actions, records, groups, pending, solo, unmatched,
    suggestions, done}。done = 无未决发票/支付，或本轮重读没有任何字段改善（提前止损）。
    """
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    try:
        session = staging.session_dir(payload.get("session_id"), request.user.id)
        records = _parse_records(payload.get("records"))
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)

    groups, pending, solo, unmatched, by_id = _pair_records(records)
    targets = _unresolved_ids(pending, solo, unmatched, by_id)[:AGENT_MAX_REREADS]
    actions, changed = [], False
    for rid in targets:
        record = by_id[rid]
        label = record.get("filename") or rid
        try:
            path = staging.file_path(session, record.get("file_id"))
        except ValueError as exc:
            return JsonResponse({"error": str(exc)}, status=400)
        candidates = [
            pair_summary(other) for other in records
            if other["id"] != rid and other.get("kind") in _AGENT_UNRESOLVED_KINDS
        ][:20]
        hints = json.dumps({
            "本文件名": record.get("filename"),
            "已知信息": pair_summary(record),
            "其他未配对记录": candidates,
        }, ensure_ascii=False)
        updated = vision.reread(path.read_bytes(), staging.file_label(path), hints)
        if not updated:
            actions.append({"file": label, "text": "重读失败，保留原识别结果"})
            continue
        if updated.get("kind") == Attachment.KIND_INVOICE and updated.get("invoice_amount") is not None:
            updated["amount"] = updated["invoice_amount"]  # 发票口径金额统一存 amount
        # 重读结果是原始 JSON（金额常为 float/数字串），先归一 Decimal 再入 records，
        # 否则第二轮配对 Decimal 与 float 混算直接 500
        updated["amount"] = decimal_or_none(updated.get("amount"))
        diffs = []
        for key in _AGENT_FIELDS:
            new = updated.get(key)
            if new not in (None, "") and new != record.get(key):
                diffs.append(f"{key}：{record.get(key) or '空'} → {new}")
                record[key] = new
        if diffs:
            changed = True
            actions.append({"file": label, "text": f"重读修正 {len(diffs)} 处：" + "；".join(diffs)})
        else:
            actions.append({"file": label, "text": "重读无修正，字段维持原识别"})
    if changed:
        groups, pending, solo, unmatched, by_id = _pair_records(records)
    still_open = _unresolved_ids(pending, solo, unmatched, by_id)
    return JsonResponse({
        "actions": actions,
        "records": _json_records(records),
        "groups": groups,
        "pending": pending,
        "solo": solo,
        "unmatched": unmatched,
        "suggestions": _field_suggestions(groups, by_id),
        "done": not changed or not still_open,
    })


@require_POST
@login_required
def batch_submit(request):
    """按确认后的分组计划一次性创建待审核条目；文件取自配对会话暂存（ADR-0008）。

    单边发票组放行：无支付记录时实付款强制 0.00（条目自带缺支付警告，审核不通过后
    学生从看板补传支付记录）；单边支付组仍整体拒绝。任一校验失败不建任何条目。
    """
    try:
        payload = json.loads(request.body or "{}")
    except ValueError:
        return JsonResponse({"error": "请求体不是合法 JSON"}, status=400)
    raw_groups = payload.get("groups") if isinstance(payload, dict) else None
    if not isinstance(raw_groups, list) or not raw_groups:
        return JsonResponse({"error": "缺少分组计划"}, status=400)
    try:
        session = staging.session_dir(payload.get("session_id"), request.user.id)
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    batch = Batch.objects.first()
    if batch is None:
        return JsonResponse({"error": "请先创建批次"}, status=400)
    if batch.archived:
        return JsonResponse({"error": "最新批次已存档，请新建批次后再提交"}, status=400)

    prepared, opened = [], []
    for number, group in enumerate(raw_groups, start=1):
        if not isinstance(group, dict):
            return JsonResponse({"error": f"第 {number} 组信息无效"}, status=400)
        title = str(group.get("title") or "").strip()[:200]
        if not title:
            return JsonResponse({"error": f"第 {number} 组缺少明细"}, status=400)
        category = Category.objects.filter(pk=group.get("category")).first()
        if category is None:
            return JsonResponse({"error": f"第 {number} 组类别无效"}, status=400)
        try:
            actual_amount = plan_decimal(group.get("actual_amount"), "实付金额")
            invoice_amount = plan_decimal(group.get("invoice_amount"), "发票金额")
        except ValueError as exc:
            return JsonResponse({"error": f"第 {number} 组{exc}"}, status=400)
        entries = group.get("files")
        if not isinstance(entries, list) or not entries:
            return JsonResponse({"error": f"第 {number} 组缺少文件"}, status=400)
        attachments, kinds = [], set()
        for entry in entries:
            if not isinstance(entry, dict):
                return JsonResponse({"error": f"第 {number} 组文件信息无效"}, status=400)
            kind = entry.get("kind")
            if kind not in KIND_LABELS:
                return JsonResponse({"error": f"第 {number} 组含未识别类型的文件，请先修正类型"}, status=400)
            try:
                path = staging.file_path(session, entry.get("file_id"))
            except ValueError as exc:
                return JsonResponse({"error": f"第 {number} 组{exc}"}, status=400)
            name = staging.file_label(path)
            error = validate_upload(SimpleNamespace(name=name, size=path.stat().st_size))
            if error:
                return JsonResponse({"error": f"{name}：{error}"}, status=400)
            handle = path.open("rb")
            opened.append(handle)
            try:
                attachments.append(build_attachment(kind, File(handle, name=name), entry))
            except ValueError as exc:
                for opened_handle in opened:
                    opened_handle.close()
                return JsonResponse({"error": f"{name}：{exc}"}, status=400)
            kinds.add(kind)
        if Attachment.KIND_INVOICE not in kinds:
            for opened_handle in opened:
                opened_handle.close()
            return JsonResponse({"error": f"第 {number} 组缺少发票，无法报销"}, status=400)
        if Attachment.KIND_PAYMENT not in kinds:
            actual_amount = Decimal("0.00")  # 单边发票组：等补传支付记录后再修订
        elif actual_amount is None:
            for opened_handle in opened:
                opened_handle.close()
            return JsonResponse({"error": f"第 {number} 组缺少实付金额"}, status=400)
        prepared.append((title, category, actual_amount, invoice_amount, attachments))

    try:
        with transaction.atomic():
            for title, category, actual_amount, invoice_amount, attachments in prepared:
                item = Item(
                    owner=request.user,
                    batch=batch,
                    position=Item.next_position(),
                    title=title,
                    category=category,
                    actual_amount=actual_amount,
                    invoice_amount=invoice_amount,
                    status=Item.STATUS_PENDING,
                )
                mark_attachments(item, [attachment_entry(a) for a in attachments], [])
                item.save()
                for attachment in attachments:
                    attachment.item = item
                    attachment.save()
                for warning in check_item(item):
                    messages.warning(request, f"{title}：{warning}")
    finally:
        for opened_handle in opened:
            opened_handle.close()
    staging.cleanup(payload.get("session_id"))
    messages.success(request, f"已批量提交 {len(prepared)} 条，等待管理员审核。")
    return redirect("board")


def _reject_archived(item):
    """存档批次只读：任何写操作一律拒绝（学生与 staff 一致），解锁走 admin 改标记。"""
    if item.batch.archived:
        raise PermissionDenied("已存档批次不可修改")

def _owned_attachment(request, pk):
    attachment = get_object_or_404(Attachment, pk=pk)
    if attachment.item.owner_id != request.user.id and not request.user.is_staff:
        raise PermissionDenied("只能修改自己的条目")
    if attachment.item.batch.archived:
        raise PermissionDenied("已存档批次不可修改")
    return attachment


def _collect_new_attachments(request):
    """从 POST 的文件与同名前缀字段构造未保存的 Attachment；返回 (附件, 错误文案)。"""
    attachments, errors = [], []
    for field, prefix, kind in UPLOAD_FIELDS:
        for index, upload in enumerate(request.FILES.getlist(field)):
            error = validate_upload(upload)
            if error:
                errors.append(f"{upload.name}：{error}")
                continue
            try:
                attachments.append(_new_attachment(request, upload, prefix, index, kind))
            except ValueError as exc:
                errors.append(f"{upload.name}：{exc}")
    return attachments, errors


def _new_attachment(request, upload, prefix, index, kind):
    def value(name):
        return (request.POST.get(f"{prefix}_{index}_{name}") or "").strip()

    return build_attachment(kind, upload, {
        "amount": value("amount"),
        "order_no": value("order_no"),
        "merchant_no": value("merchant_no"),
        "invoice_no": value("invoice_no"),
        "ocr": value("ocr"),
    })


def _staff(user):
    return user.is_staff


def _current_batch(request):
    batch_id = request.GET.get("batch", "")
    if batch_id.isdigit():
        batch = Batch.objects.filter(pk=batch_id).first()
        if batch is not None:
            return batch
    return Batch.objects.first()


def _approved_by_category(batch):
    """当前批次 approved 条目，按类别顺序 + 持久化顺序排序，返回 [(category, items)]。"""
    approved = (
        Item.objects.filter(status=Item.STATUS_APPROVED, batch=batch)
        .select_related("owner", "category")
        .prefetch_related("attachments")
        .order_by("category__order", "category__id", "position", "id")
    )
    return [
        (category, list(group))
        for category, group in groupby(approved, key=lambda item: item.category)
    ]


def _download_filename(prefix, batch, extension):
    """批次名只去除会破坏 Content-Disposition 的字符（引号/换行），中文原样保留。"""
    name = (batch.name if batch else "").replace('"', "").replace("\r", "").replace("\n", "")
    stem = f"{prefix}_{name}" if name else prefix
    return f'attachment; filename="{stem}.{extension}"'


@user_passes_test(_staff)
def export_excel(request):
    batch = _current_batch(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "报销明细"
    header = ["序号", "明细", "类别", "实付款", "付款人", "支付记录", "票据"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    seq = 0
    subtotal_rows = []
    for category, items in _approved_by_category(batch):
        start = None
        for item in items:
            seq += 1
            ws.append([
                seq,
                item.title,
                category.name,
                float(item.actual_amount),
                item.payer_name,
                "已整理",
                "已整理",
            ])
            row = ws.max_row
            amount_cell = ws.cell(row=row, column=4)
            amount_cell.number_format = AMOUNT_FORMAT
            # 多张附件时超链接指向首张，完整材料走 zip 导出（ADR-0005）
            if item.first_payment is not None:
                ws.cell(row=row, column=6).hyperlink = request.build_absolute_uri(
                    reverse("file_serve", args=[item.first_payment.pk])
                )
            if item.first_invoice is not None:
                ws.cell(row=row, column=7).hyperlink = request.build_absolute_uri(
                    reverse("file_serve", args=[item.first_invoice.pk])
                )
            if start is None:
                start = row
        ws.append(["小记", None, None, f"=SUM(D{start}:D{ws.max_row})", None, None, None])
        subtotal_row = ws.max_row
        ws.cell(row=subtotal_row, column=4).number_format = AMOUNT_FORMAT
        subtotal_rows.append(subtotal_row)

    if subtotal_rows:
        ws.append(["合计", None, None, "=" + "+".join(f"D{r}" for r in subtotal_rows), None, None, None])
        ws.cell(row=ws.max_row, column=4).number_format = AMOUNT_FORMAT

    widths = {"A": 6, "B": 40, "C": 12, "D": 10, "E": 10, "F": 10, "G": 10}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = _download_filename("报销汇总", batch, "xlsx")
    return response


@user_passes_test(_staff)
def export_zip(request):
    batch = _current_batch(request)
    buffer = BytesIO()
    seen = set()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        seq = 0
        for category, items in _approved_by_category(batch):
            for item in items:
                seq += 1  # 与 Excel"序号"列一致（全局连续）
                counters = {}
                payments = [a for a in item.attachments.all()
                            if a.kind == Attachment.KIND_PAYMENT]
                first_invoice = item.first_invoice
                # 各支付记录的平台单号逐行写在第一张发票附件首页（CONTEXT.md 票据标注）
                order_lines = stamping.order_lines(payments) if first_invoice else []
                for attachment in item.attachments.all():
                    kind_label = KIND_LABELS[attachment.kind]
                    counters[kind_label] = counters.get(kind_label, 0) + 1
                    data = attachment.file.read()
                    extension = attachment.file.name.rsplit(".", 1)[-1]
                    lines = (order_lines if first_invoice
                             and attachment.pk == first_invoice.pk else [])
                    stamped = stamping.stamped_pdf(data, attachment.file.name, f"{seq:02d}", lines)
                    if stamped is not None:
                        data, extension = stamped, "pdf"
                    name = (
                        f"{category.name}_{seq:02d}_{item.payer_name}"
                        f"_{kind_label}_{counters[kind_label]}.{extension}"
                    )
                    if name in seen:
                        name = (
                            f"{category.name}_{seq:02d}_{item.payer_name}"
                            f"_{kind_label}_{counters[kind_label]}_{item.pk}.{extension}"
                        )
                    seen.add(name)
                    zf.writestr(name, data)
    response = HttpResponse(buffer.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = _download_filename("报销材料", batch, "zip")
    return response


def session_check(request):
    """nginx auth_request 子请求（跨子域名共享 session，如 mineru.cislunarspace.cn）。

    已登录返回 204，未登录返回 401；由边缘 nginx 将 401 重定向到统一登录页。
    """
    if request.user.is_authenticated:
        return HttpResponse(status=204)
    return HttpResponse(status=401)
